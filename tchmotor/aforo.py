"""Aforo físico (medición en campo) y aforo sintético (histórico).

El aforo físico replica la fórmula usada en campo por Riopaila
(turnos-jornales/lib/eval/aforo.ts): TCH = población × peso × (10/espaciamiento),
con dos métodos (pesos medidos y peso calculado del diámetro).

El aforo sintético es la mediana histórica de esos parámetros por
variedad × soca — disponible para toda suerte aunque no se afore hoy.
"""
from __future__ import annotations
import math
from statistics import median
from typing import Iterable

# Densidad de la caña ≈ 1.04 g/cm³ (kg/cm³). Igual que el app de campo.
DENSIDAD_CANA = 0.00104


def _soca_ok(s) -> bool:
    """Soca utilizable: no None y no NaN (NaN != NaN)."""
    return s is not None and s == s


def peso_por_diametro(diametro_cm: float, longitud_cm: float) -> float:
    """Peso (kg) de un tallo estimado de su geometría (cilindro)."""
    return (math.pi / 4) * diametro_cm * diametro_cm * longitud_cm * DENSIDAD_CANA


def tch_por_pesos(poblacion: float, peso_kg: float, espaciamiento_m: float) -> float:
    """TCH por el método de pesos: población(tallos/m) × peso(kg/tallo) × 10/esp.

    `espaciamiento` es la distancia entre surcos (m). Devuelve 0 si es <= 0
    (dato incompleto), igual que el app de campo, para no dividir por cero.
    """
    if espaciamiento_m <= 0:
        return 0.0
    return poblacion * peso_kg * (10.0 / espaciamiento_m)


def tch_por_diametro(
    poblacion: float, diametro_cm: float, longitud_cm: float, espaciamiento_m: float
) -> float:
    """TCH por el método del diámetro (usa el peso calculado de la geometría)."""
    return tch_por_pesos(
        poblacion, peso_por_diametro(diametro_cm, longitud_cm), espaciamiento_m
    )


class AforoSintetico:
    """Estadísticas históricas de tallos por variedad × soca, con respaldo.

    Se construye una vez con los registros de la base de aforos y responde,
    para cualquier (variedad, soca), la población/peso/diámetro típicos y el
    TCH sintético (mediana del TCH calculado). El respaldo es
    variedad×soca → variedad → global, para no quedar sin dato.
    """

    def __init__(self, registros: Iterable[dict], min_n: int = 4):
        # registros: dicts con variedad, soca, pob, peso, dia, tch (calculado)
        regs = [r for r in registros if r.get("pob") is not None]
        self._vs: dict[tuple, dict] = {}
        self._v: dict[str, dict] = {}
        self._global = self._resumen(regs)
        by_vs: dict[tuple, list] = {}
        by_v: dict[str, list] = {}
        for r in regs:
            v = r.get("variedad")
            s = r.get("soca")
            by_v.setdefault(v, []).append(r)
            if _soca_ok(s):
                by_vs.setdefault((v, round(s)), []).append(r)
        self._vs = {k: self._resumen(g) for k, g in by_vs.items() if len(g) >= min_n}
        self._v = {k: self._resumen(g) for k, g in by_v.items() if len(g) >= min_n}

    @staticmethod
    def _med(regs: list, campo: str) -> float | None:
        vals = [r[campo] for r in regs if r.get(campo) is not None]
        return float(median(vals)) if vals else None

    def _resumen(self, regs: list) -> dict:
        return {
            "pob": self._med(regs, "pob"),
            "peso": self._med(regs, "peso"),
            "dia": self._med(regs, "dia"),
            "aforo": self._med(regs, "tch"),
            "n": len(regs),
        }

    def filas(self) -> list[dict]:
        """Filas para persistir en `tch_aforo_sintetico` (soca=-1 = agregado de
        variedad, el respaldo). Excluye el global, que se recalcula en memoria."""
        out = []
        for (var, soca), r in self._vs.items():
            out.append(dict(variedad=var, soca=int(soca), poblacion=r["pob"],
                            peso=r["peso"], diametro=r["dia"], aforo_sint=r["aforo"], n=r["n"]))
        for var, r in self._v.items():
            out.append(dict(variedad=var, soca=-1, poblacion=r["pob"],
                            peso=r["peso"], diametro=r["dia"], aforo_sint=r["aforo"], n=r["n"]))
        return out

    def lookup(self, variedad: str, soca) -> dict:
        """Población/peso/diámetro/aforo sintético para esa variedad×soca."""
        if _soca_ok(soca) and (variedad, round(soca)) in self._vs:
            return self._vs[(variedad, round(soca))]
        if variedad in self._v:
            return self._v[variedad]
        return self._global


def _num(x):
    try:
        return float(str(x).replace(",", "."))
    except Exception:
        return float("nan")


def registros_aforo_excel(ruta: str) -> list[dict]:
    """Registros de tallos (variedad, soca, población, peso, diámetro, TCH calculado) de la base
    de aforos de Agronomía ("Base de aforos completa 2026.xlsx"), para `AforoSintetico`."""
    import openpyxl

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    out = []
    for name in ["Aforos 2024-2026", "GENERAL DE LA SUERTE", "BASE 2022-2023", "AFOROS 2022-2023"]:
        if name not in wb.sheetnames:
            continue
        rows = list(wb[name].iter_rows(values_only=True, max_col=34))
        hi = next((i for i, r in enumerate(rows[:6]) if any(str(c).strip().upper() == "VARIEDAD" for c in r if c)), None)
        if hi is None:
            continue
        H = {str(c).strip().lower(): j for j, c in enumerate(rows[hi]) if c}

        def fnd(*s):
            return next((j for k, j in H.items() if all(x in k for x in s)), None)

        cv, cc, cpob, ctch = fnd("variedad"), fnd("corte"), fnd("tallos", "mt"), fnd("tch", "calcul")
        cpe, cdi = fnd("total", "peso"), fnd("total", "diam")
        if cv is None or cpob is None:
            continue
        for r in rows[hi + 1:]:
            var = str(r[cv]).strip() if r[cv] else None
            pob = _num(r[cpob])
            if not var or not (var.startswith("CC") or var == "VARIAS") or not (3 < pob < 25):
                continue
            out.append(dict(variedad=var, soca=_num(r[cc]) if cc is not None else None,
                            pob=pob, peso=_num(r[cpe]) if cpe else None,
                            dia=_num(r[cdi]) if cdi else None,
                            tch=_num(r[ctch]) if ctch else None))
    wb.close()
    return out


class AforoSinteticoTabla:
    """Aforo sintético desde las filas ya agregadas de `tch_aforo_sintetico`
    (Supabase), en vez de recomputarlo de los aforos crudos. Mismo `lookup`.
    """

    def __init__(self, filas: Iterable[dict]):
        self._vs: dict[tuple, dict] = {}
        self._v: dict[str, dict] = {}
        for f in filas:
            d = {"pob": f.get("poblacion"), "peso": f.get("peso"),
                 "dia": f.get("diametro"), "aforo": f.get("aforo_sint")}
            if int(f.get("soca", -1)) == -1:
                self._v[f["variedad"]] = d
            else:
                self._vs[(f["variedad"], int(f["soca"]))] = d
        vs = list(self._v.values()) or list(self._vs.values())
        self._global = {
            k: (float(median([r[k] for r in vs if r.get(k) is not None]))
                if any(r.get(k) is not None for r in vs) else None)
            for k in ("pob", "peso", "dia", "aforo")
        }

    def lookup(self, variedad: str, soca) -> dict:
        if _soca_ok(soca) and (variedad, round(soca)) in self._vs:
            return self._vs[(variedad, round(soca))]
        if variedad in self._v:
            return self._v[variedad]
        return self._global
