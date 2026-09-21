# -*- coding: utf-8 -*-
"""Carga el libro de aforos de Agronomía a la tabla `tch_aforo_vivo` (Supabase de
Rio Map), que es el feed del aforo en vivo para la FUSIÓN.

`aforo_vivo.aforos_vigentes` toma de ahí, por suerte, el aforo más reciente antes
de la corrida y devuelve el PROMEDIO de los métodos (diámetro + pesos), validado
como el mejor. Idempotente: upsert por (suerte, date).

Requisito previo: aplicar sql/0003_tch_aforo_vivo.sql en el SQL Editor de Rio Map.

Uso:
  python scripts/cargar_aforos.py --dry
  python scripts/cargar_aforos.py --libro "C:\\Users\\Agr349\\Documents\\aforo con diametro.xlsx"
"""
import argparse, sys
import openpyxl

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
sys.path.insert(0, r"C:\Users\Agr349\tch-motor\scripts")
from tchmotor import Supabase
import correr_estimados as ce

LIBRO = r"C:\Users\Agr349\Documents\aforo con diametro.xlsx"
# columnas del libro (0-based): Sec-Ste=0, Fecha Evaluación=3, Corte=5, Variedad=8,
# Diámetro(cm)=9, Población=11, TCH DIÁMETRO=13, TCH Pesos=15
C_SEC, C_FEVAL, C_VAR, C_DIA_CM, C_POB, C_TCH_DIA, C_TCH_PES = 0, 3, 8, 9, 11, 13, 15


def num(x):
    try:
        return float(x)
    except Exception:
        return None


def fecha(x):
    if hasattr(x, "date"):
        return x.date().isoformat()
    s = str(x)[:10]
    return s if len(s) == 10 and s[4] == "-" else None


def leer_libro(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas, vistos = [], set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        sec = row[C_SEC]
        if not sec:
            continue
        sec = str(sec).strip()
        d = fecha(row[C_FEVAL])
        dia = num(row[C_TCH_DIA]); pes = num(row[C_TCH_PES])
        if not d or (dia is None and pes is None):
            continue
        k = (sec, d)
        if k in vistos:      # dedupe dentro del libro (misma suerte+fecha)
            continue
        vistos.add(k)
        filas.append({"suerte": sec, "date": d,
                      "tch_diametro": round(dia, 2) if dia is not None else None,
                      "tch_pesos": round(pes, 2) if pes is not None else None,
                      "variedad": str(row[C_VAR]).strip() if row[C_VAR] else None,
                      "poblacion": num(row[C_POB]), "diametro": num(row[C_DIA_CM]),
                      "fuente": "agronomia_xlsx"})
    wb.close()
    return filas


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--libro", default=LIBRO)
    ap.add_argument("--dry", action="store_true")
    args = ap.parse_args()

    filas = leer_libro(args.libro)
    print(f"aforos en el libro (con suerte+fecha+algún TCH): {len(filas)}")
    if args.dry:
        for f in filas[:4]:
            print("  ", f)
        print("[dry] no se escribió. Recuerda aplicar sql/0003_tch_aforo_vivo.sql antes de cargar.")
        return
    if not filas:
        print("nada que cargar.")
        return

    url = ce._env(r"C:\Users\Agr349\agrocontrol-campo\.env.local", "NEXT_PUBLIC_SUPABASE_URL")
    svc = ce._env(r"C:\Users\Agr349\turnos-jornales\.env.local", "RIOMAP_SERVICE_ROLE_KEY")
    sb = Supabase(url=url, key=svc)
    n = 0
    for i in range(0, len(filas), 200):
        n += sb.upsert("tch_aforo_vivo", filas[i:i + 200], on_conflict="suerte,date")
    print(f"cargados/actualizados {n} aforos en tch_aforo_vivo. "
          f"La fusión los usará en la próxima corrida (poblar).")


if __name__ == "__main__":
    main()
