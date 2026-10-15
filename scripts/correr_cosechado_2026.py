"""Corre el modelo sobre lo COSECHADO en 2026, out-of-sample.

Para cada suerte cortada en 2026 (del zqm 26) reconstruye la ventana satelital de
la soca que se cortó (fecha de corte - edad) y predice el TCH con el motor entrenado
solo con 2019-2025 (2026 no está en el dataset → naturalmente out-of-sample).
Así la comparación es manzana con manzana: mismo ciclo, misma soca.

Salida: CSV incremental + Excel con 'Por suerte' y 'Resumen'.
Uso: python scripts/correr_cosechado_2026.py [--lote 40]
"""
import argparse, csv, os, sys, json
from datetime import date, timedelta
from concurrent.futures import ThreadPoolExecutor
import numpy as np, pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
sys.path.insert(0, r"C:\Users\Agr349\tch-motor\scripts")
from tchmotor import (Supabase, AforoSinteticoTabla, build_features, MotorTCH,
                      curva_suerte, map_variedad)
import correr_estimados as ce

ZQM26 = r"C:\Users\Agr349\Documents\zqm 26.xlsx"
MAESTRO = r"C:\Users\Agr349\agrocontrol-campo\public\data"
IDX = r"C:\Users\Agr349\tch-motor\data\tablones_index.json"
DS = r"C:\Users\Agr349\tch-motor\data\dataset_entrenamiento.jsonl"
CSVOUT = r"C:\Users\Agr349\Documents\Cosechado_2026_estimado.csv"
XLSX = r"C:\Users\Agr349\Documents\Cosechado_2026_estimado.xlsx"
# columnas zqm 26 (0-based): Mes=3, sec=2, Corte=14, Var=15, Area=16, Edad=17,
# TC-Total=19, TCH=20, Planta=45. header en fila 10, datos desde fila 11.
C_SEC, C_MES, C_CORTE, C_VAR, C_AREA, C_EDAD, C_TCTOT, C_TCH, C_PLANTA = 2, 3, 14, 15, 16, 17, 19, 20, 45
_DIAS = 30.44
MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def num(x):
    try:
        return float(str(x).replace(",", "."))
    except Exception:
        return np.nan


def cosechadas():
    wb = openpyxl.load_workbook(ZQM26, read_only=True, data_only=True)
    ws = wb["ZQM101"]; best = {}
    for row in ws.iter_rows(min_row=11, values_only=True):
        sec = row[C_SEC]
        tch, tct, area, edad = num(row[C_TCH]), num(row[C_TCTOT]), num(row[C_AREA]), num(row[C_EDAD])
        if not sec or np.isnan(tch) or np.isnan(area) or area <= 0 or np.isnan(edad):
            continue
        sec = str(sec).strip()
        planta = str(row[C_PLANTA]).strip().lower() if row[C_PLANTA] else ""
        planta = "castilla" if planta.startswith("cas") else ("riopaila" if planta.startswith("rio") else planta)
        if sec not in best or (tct or 0) > best[sec]["_t"]:
            best[sec] = {"tch_real": tch, "area": area, "edad": edad, "mes": num(row[C_MES]),
                         "corte": num(row[C_CORTE]), "var": str(row[C_VAR]).strip() if row[C_VAR] else None,
                         "planta_zqm": planta, "_t": tct or 0}
    wb.close()
    return best


def construir_excel():
    df = pd.read_csv(CSVOUT, encoding="utf-8-sig")
    df = df.sort_values(["planta", "sec_ste"]).reset_index(drop=True)
    # resumen agregado ponderado por area
    def agg(g):
        area = g["area_ha"].sum()
        est = (g["tch_estimado"] * g["area_ha"]).sum() / area
        real = (g["tch_real"] * g["area_ha"]).sum() / area
        mae = (g["tch_estimado"] - g["tch_real"]).abs().mean()
        mape = ((g["tch_estimado"] - g["tch_real"]).abs() / g["tch_real"]).mean() * 100
        d20 = (((g["tch_estimado"] - g["tch_real"]).abs() / g["tch_real"]) <= 0.20).mean() * 100
        return pd.Series({"suertes": len(g), "area_ha": round(area), "tch_est_pond": round(est, 1),
                          "tch_real_pond": round(real, 1), "desfase_%": round((est / real - 1) * 100, 1),
                          "MAE_tha": round(mae, 1), "MAPE_%": round(mape, 1), "dentro_20%": round(d20)})
    res = df.groupby("planta").apply(agg).reset_index()
    tot = agg(df); tot["planta"] = "TOTAL"
    res = pd.concat([res, pd.DataFrame([tot])[res.columns]], ignore_index=True)

    with pd.ExcelWriter(XLSX, engine="openpyxl") as xl:
        df.to_excel(xl, sheet_name="Por suerte", index=False)
        res.to_excel(xl, sheet_name="Resumen", index=False)
    wb = openpyxl.load_workbook(XLSX)
    hf = PatternFill("solid", fgColor="1F5A3D"); hfont = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="E5E5E5")
    for ws in wb.worksheets:
        for c in ws[1]:
            c.fill = hf; c.font = hfont; c.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            wmax = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(wmax + 2, 10), 22)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = Border(bottom=thin)
    wb.save(XLSX)
    print(f"\nExcel -> {XLSX}")
    print("Resumen:")
    for _, r in res.iterrows():
        print(f"  {r['planta']:>9}: {int(r['suertes']):>4} suertes | est {r['tch_est_pond']} vs real "
              f"{r['tch_real_pond']} t/ha | desfase {r['desfase_%']:+}% | MAE {r['MAE_tha']} | "
              f"MAPE {r['MAPE_%']}% | dentro 20% {int(r['dentro_20%'])}%")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lote", type=int, default=48)
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--solo-excel", action="store_true")
    args = ap.parse_args()
    if args.solo_excel:
        construir_excel(); return

    os.environ.setdefault("TCH_SUPABASE_URL", ce._env(r"C:\Users\Agr349\agrocontrol-campo\.env.local", "NEXT_PUBLIC_SUPABASE_URL"))
    os.environ.setdefault("TCH_SUPABASE_KEY", ce._env(r"C:\Users\Agr349\agrocontrol-campo\.env.local", "NEXT_PUBLIC_SUPABASE_ANON_KEY"))
    sb = Supabase()
    vm = sb.cargar_var_map()
    synth = AforoSinteticoTabla(sb.cargar_aforo_sintetico())

    maestro = {}
    for fn in ("maestro_suertes.json", "maestro_castilla.json"):
        maestro.update(json.load(open(os.path.join(MAESTRO, fn), encoding="utf-8")))
    idx = json.load(open(IDX, encoding="utf-8"))

    # pool de entrenamiento completo 2019-2025 (2026 no está aqui -> out-of-sample)
    Xrows, ys = [], []
    for line in open(DS, encoding="utf-8"):
        r = json.loads(line)
        if not r["curves"].get("NDVI"):
            continue
        f = build_features(r, synth, vm)
        if f.get("ndvi_n", 0) < 3:
            continue
        Xrows.append(f); ys.append(r["tch"])
    Xpool = pd.DataFrame(Xrows); ypool = np.array(ys)
    print(f"pool de entrenamiento: {len(ypool)} eventos (2019-2025)", flush=True)
    motor = MotorTCH(n_estimators=600).entrenar(Xpool, ypool)

    cos = cosechadas()
    suertes = []
    sin_idx = 0
    for sec, c in cos.items():
        if c["var"] in (None, "V19", "VARIAS", "RENOVACION"):
            continue
        if sec not in idx:
            sin_idx += 1; continue
        tab_id, _h, planta = idx[sec]
        mes = int(c["mes"]) if c["mes"] and not np.isnan(c["mes"]) else 9
        harvest = date(2026, min(max(mes, 1), 12), 15)
        cut = (harvest - timedelta(days=float(c["edad"]) * _DIAS)).isoformat()
        zona = (maestro.get(sec) or {}).get("zona_agroecologica")
        suertes.append({"sec": sec, "planta": planta, "planta_zqm": c["planta_zqm"], "tab_id": tab_id,
                        "cut": cut, "harvest": harvest.isoformat(), "mes": mes,
                        "edad": round(float(c["edad"]), 1),
                        "corte": int(c["corte"]) if c["corte"] and not np.isnan(c["corte"]) else None,
                        "var": c["var"], "zona": zona, "area": c["area"], "real": c["tch_real"]})
    print(f"cosechadas 2026: {len(cos)} | con tablon satelital: {len(suertes)} | sin index: {sin_idx}", flush=True)

    cols = ["sec_ste", "planta", "variedad", "soca", "edad_meses", "mes_corte", "area_ha",
            "tch_estimado", "tch_std", "tch_real", "error_tha", "error_pct", "modo", "fuente"]
    # reanudar: si ya hay CSV con filas, no repetir esas suertes
    done = set()
    if os.path.exists(CSVOUT):
        try:
            prev = pd.read_csv(CSVOUT, encoding="utf-8-sig")
            done = set(prev["sec_ste"].astype(str))
        except Exception:
            done = set()
    if done:
        suertes = [s for s in suertes if s["sec"] not in done]
        fh = open(CSVOUT, "a", encoding="utf-8-sig", newline="")
        w = csv.DictWriter(fh, fieldnames=cols)
        print(f"reanudando: {len(done)} ya hechas, faltan {len(suertes)}", flush=True)
    else:
        fh = open(CSVOUT, "w", encoding="utf-8-sig", newline="")
        w = csv.DictWriter(fh, fieldnames=cols); w.writeheader(); fh.flush()

    def bajar(s):
        try:
            return s, curva_suerte(s["tab_id"], s["planta"], s["cut"], s["harvest"])
        except Exception:
            return s, None

    hechas = 0; sin_nube = 0
    pool = ThreadPoolExecutor(max_workers=args.workers)
    for i in range(0, len(suertes), args.lote):
        lote = suertes[i:i + args.lote]
        feats, meta = [], []
        for s, curves in pool.map(bajar, lote):
            if not curves:
                sin_nube += 1; continue
            f = build_features({"variedad": s["var"], "corte": s["corte"], "edad": s["edad"],
                                "zona_agro": s["zona"], "ha": s["area"], "mes": s["mes"],
                                "cut": s["cut"], "curves": curves}, synth, vm)
            if f.get("ndvi_n", 0) < 3:
                sin_nube += 1; continue
            feats.append(f); meta.append(s)
        if feats:
            mu, sd = motor.predecir_con_banda(pd.DataFrame(feats))
            for s, e, band in zip(meta, mu, sd):
                e = round(float(e), 1); r = s["real"]
                w.writerow({"sec_ste": s["sec"], "planta": s["planta_zqm"] or s["planta"],
                            "variedad": map_variedad(s["var"], vm), "soca": s["corte"],
                            "edad_meses": s["edad"], "mes_corte": MESES[s["mes"]], "area_ha": round(s["area"], 2),
                            "tch_estimado": e, "tch_std": round(float(band), 1), "tch_real": round(r, 1),
                            "error_tha": round(e - r, 1), "error_pct": round((e - r) / r * 100, 1),
                            "modo": "motor", "fuente": "out-of-sample 2026"})
                hechas += 1
            fh.flush()
        print(f"  [{min(i + args.lote, len(suertes))}/{len(suertes)}] estimadas {hechas} | sin nube {sin_nube}", flush=True)
    fh.close()
    print(f"\nlisto: {hechas} estimadas de {len(suertes)} con tablon ({sin_nube} sin senal satelital)", flush=True)
    construir_excel()


if __name__ == "__main__":
    main()
