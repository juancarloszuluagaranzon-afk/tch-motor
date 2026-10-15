# -*- coding: utf-8 -*-
"""Aforo de Agronomía frente al motor en las suertes que FALTAN por cosechar (resto del año y el siguiente). Riopaila.

Sin TCH real todavía: la comparación es aforo contra motor, y al lado de cada grupo va cómo les fue a los dos
contra el real en las cosechas ya cerradas (por defecto, desde el 1 de septiembre), para saber en cuál confiar
cuando no coinciden.

Suertes: maestro oficial, uso CAÑA, sin renovación, con próximo corte en el año en curso o el siguiente.
Motor: estimados.json de la última corrida, del MISMO corte que va a cosecharse (ciclo en curso).
Aforo: «Promedio» del libro «Aforos con Diámetro» de Agronomía, último aforo del ciclo en curso
       (mismo número de corte y evaluado después del último corte o la siembra).
Los aforos se hacen hacia los 10 meses: una suerte entra a la comparación cuando ya la aforaron.

Uso:  .venv-gee\\Scripts\\python.exe scripts\\comparar_aforo_motor_pendientes.py [--referencia-desde 2026-09-01] [--salida x.xlsx]
"""
import argparse
import json
import os
import shutil
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from tchmotor import rutas  # noqa: E402
from presupuesto_cosecha import leer_maestro  # noqa: E402
import comparar_real_motor_aforo as cra  # noqa: E402

MESES = cra.MESES
DIAS = 30.4375
BANDAS = [("Total", 1, 99), ("Plantilla (corte 1)", 1, 1), ("Socas (corte 2 o más)", 2, 99), ("   Soca 2-3", 2, 3),
          ("   Soca 4-6", 4, 6), ("   Soca 7-10", 7, 10), ("   Soca 11 o más", 11, 99)]


def leer_aforos():
    src = rutas.AFOROS_DIAMETRO
    with tempfile.TemporaryDirectory() as tmp:  # OneDrive bloquea la lectura directa
        copia = Path(tmp) / "aforos.xlsx"
        shutil.copy2(src, copia)
        a = pd.read_excel(copia, "Aforos", header=1)
    a = a.rename(columns={a.columns[0]: "sec", a.columns[3]: "f_eval", a.columns[5]: "corte_af", a.columns[7]: "edad_af",
                          a.columns[13]: "af_diam", a.columns[14]: "af_diam12", a.columns[15]: "af_pesos", a.columns[16]: "aforo"})
    a["sec"] = a.sec.astype(str).str.strip()
    a["f_eval"] = pd.to_datetime(a.f_eval, errors="coerce")
    return a


def armar(hoy):
    ma, fuente_maestro = leer_maestro("riopaila")
    a = leer_aforos()
    est = {e["sec_ste"]: e for e in json.load(open(rutas.TABLERO_DATA / "estimados.json", encoding="utf-8"))
           if e["planta"] == "riopaila"}
    filas = []
    for s, v in ma.items():
        if v["var"].upper() == "RENOVACION" or not v["prox"] or not (hoy.year <= v["prox"].year <= hoy.year + 1):
            continue
        ini = max([d for d in (v["siembra"], v["ult"]) if d] or [None])
        g = a[(a.sec == s) & (a.corte_af == v["ncorte"])]
        if ini:
            g = g[g.f_eval.dt.date > ini]
        g = g.sort_values("f_eval").tail(1)
        e = est.get(s)
        motor = float(e["tch_estimado"]) if e is not None and int(e["soca"]) == v["ncorte"] else None
        af = g.iloc[0] if len(g) else None
        filas.append(dict(
            sec_ste=s, finca=v["finca"], zona=f"Zona {v['zona']}", tecnico=v["tecnico"], variedad=v["var"], corte=v["ncorte"],
            anio=v["prox"].year, mes=f"{MESES[v['prox'].month - 1]} {v['prox'].year}", prox=v["prox"],
            edad_hoy=(hoy - ini).days / DIAS if ini else None, area=v["area"], motor=motor,
            aforo=None if af is None else float(af.aforo), af_diam=None if af is None else float(af.af_diam),
            af_diam12=None if af is None else float(af.af_diam12), af_pesos=None if af is None else float(af.af_pesos),
            f_eval=None if af is None else af.f_eval.date(), edad_af=None if af is None else float(af.edad_af),
            tch_ppto=v["tch_ppto"], tch_ant=v["tch_ant"]))
    d = pd.DataFrame(filas).sort_values(["prox", "sec_ste"]).reset_index(drop=True)
    return d, fuente_maestro


def referencia(desde, hasta):
    """Cómo les fue al aforo y al motor contra el real en las cosechas ya cerradas, por grupo."""
    r = cra.armar(desde, hasta)
    t = r[r.motor.notna() & r.aforo.notna()]
    pond = lambda g, k: (g[k] * g.area_ef).sum() / g.area_ef.sum()  # noqa: E731

    def par(g):
        if len(g) < 3:
            return None
        return (pond(g, "aforo") / pond(g, "tch_real") - 1, pond(g, "motor") / pond(g, "tch_real") - 1, len(g))

    ref = {("banda", nom.strip()): par(t[t.corte.between(lo, hi)]) for nom, lo, hi in BANDAS}
    ref.update({("var", v): par(g) for v, g in t.groupby("variedad")})
    return ref, len(t)


# ======================================================================== Excel
F = "Arial"
FB, FN, FH = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=10, bold=True, color="FFFFFF")
FIN, FNOTE, FTIT = Font(name=F, size=10, color="0000FF"), Font(name=F, size=9, italic=True, color="595959"), Font(name=F, size=14, bold=True)
F9 = Font(name=F, size=9, color="595959")
HEAD, CALC, TOT = (PatternFill("solid", fgColor=c) for c in ("2F3A33", "EFEFEF", "F4F4F0"))
FMOT, FAFO, FREF = (PatternFill("solid", fgColor=c) for c in ("FBE3D6", "D5F2E7", "EDEDED"))
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
PCT = "+0.0%;-0.0%;0.0%"


def escribir(d, ref, n_ref, desde_ref, salida, hoy, fuente_maestro):
    wb = Workbook()
    ws = wb.active
    ws.title = "Detalle"
    cols = [("Suerte", 10, None), ("Hacienda", 18, None), ("Zona", 8, None), ("Técnico", 18, None), ("Variedad", 12, None),
            ("Corte", 6, "0"), ("Tipo", 9, None), ("Grupo de corte", 11, None), ("Año de cosecha", 8, "0"),
            ("Mes previsto", 14, None), ("Fecha prevista (maestro)", 11, "dd/mm/yyyy"), ("Edad hoy (meses)", 8, "0.0"),
            ("Área neta (ha)", 9, "0.00"), ("TCH motor", 9, "0.0"), ("TCH aforo (Promedio)", 10, "0.0"),
            ("Aforo Diámetro", 9, "0.0"), ("Aforo Diámetro 12 m", 9, "0.0"), ("Aforo Pesos", 9, "0.0"),
            ("Fecha del aforo", 11, "dd/mm/yyyy"), ("Edad al aforo (meses)", 8, "0.0"), ("Ref. TCH ppto (maestro)", 9, "0.0"),
            ("Ref. TCH anterior", 9, "0.0"), ("Tiene los dos", 8, None), ("Aforo frente al motor", 9, PCT),
            ("Revisar (dif. > 15 %)", 9, None),
            ("Motor × ha", 10, "#,##0"), ("Aforo × ha", 10, "#,##0"), ("Diámetro × ha", 10, "#,##0"),
            ("Diám. 12 m × ha", 10, "#,##0"), ("Pesos × ha", 10, "#,##0"), ("|Aforo − motor| ÷ motor", 9, "0.0%"),
            ("Dentro de ±10 %", 8, "0"), ("Ton motor (todas)", 10, "#,##0")]
    K = {n: get_column_letter(i + 1) for i, (n, _, _) in enumerate(cols)}
    for i, (n, w, _) in enumerate(cols, 1):
        c = ws.cell(1, i, n)
        c.alignment = CEN
        c.fill, c.font = (CALC, Font(name=F, size=9, bold=True, color="404040")) if i > 25 else (HEAD, FH)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 44
    n1 = len(d) + 1
    nan = lambda v: None if v is None or (isinstance(v, float) and v != v) else v  # noqa: E731
    azul = {"Área neta (ha)", "TCH motor", "TCH aforo (Promedio)", "Aforo Diámetro", "Aforo Diámetro 12 m", "Aforo Pesos"}
    for k, x in enumerate(d.itertuples(index=False), start=2):
        banda = "Plantilla" if x.corte == 1 else "Soca 2-3" if x.corte <= 3 else "Soca 4-6" if x.corte <= 6 else \
            "Soca 7-10" if x.corte <= 10 else "Soca 11+"
        vals = [x.sec_ste, x.finca, x.zona, x.tecnico, x.variedad, x.corte, "Plantilla" if x.corte == 1 else "Soca", banda,
                x.anio, x.mes, x.prox, nan(x.edad_hoy), x.area, nan(x.motor), nan(x.aforo), nan(x.af_diam), nan(x.af_diam12),
                nan(x.af_pesos), x.f_eval if isinstance(x.f_eval, date) else None, nan(x.edad_af), nan(x.tch_ppto), nan(x.tch_ant)]
        for j, v in enumerate(vals, 1):
            ws.cell(k, j, nan(v))
        M, O, J = (f"{K[n]}{k}" for n in ("TCH motor", "TCH aforo (Promedio)", "Área neta (ha)"))
        P, Q, R = (f"{K[n]}{k}" for n in ("Aforo Diámetro", "Aforo Diámetro 12 m", "Aforo Pesos"))
        W = f"{K['Tiene los dos']}{k}"
        f = {"Tiene los dos": f'=IF(AND(ISNUMBER({M}),ISNUMBER({O})),"Sí","No")',
             "Aforo frente al motor": f'=IF({W}="Sí",{O}/{M}-1,"")',
             "Revisar (dif. > 15 %)": f'=IF({W}="Sí",IF(ABS({O}/{M}-1)>0.15,"Revisar",""),"")',
             "Motor × ha": f"=IF(ISNUMBER({M}),{M}*{J},0)", "Aforo × ha": f"=IF(ISNUMBER({O}),{O}*{J},0)",
             "Diámetro × ha": f"=IF(ISNUMBER({P}),{P}*{J},0)", "Diám. 12 m × ha": f"=IF(ISNUMBER({Q}),{Q}*{J},0)",
             "Pesos × ha": f"=IF(ISNUMBER({R}),{R}*{J},0)",
             "|Aforo − motor| ÷ motor": f'=IF({W}="Sí",ABS({O}/{M}-1),0)',
             "Dentro de ±10 %": f'=IF({W}="Sí",IF(ABS({O}/{M}-1)<=0.1,1,0),0)',
             "Ton motor (todas)": f"=IF(ISNUMBER({M}),{M}*{J},0)"}
        for nom, frm in f.items():
            ws[f"{K[nom]}{k}"] = frm
        for j, (n, _, fmt) in enumerate(cols, 1):
            cell = ws.cell(k, j)
            cell.font = F9 if j > 25 else (FIN if n in azul and cell.value is not None else FB)
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n1}"
    ws.cell(n1 + 2, 1, "Azul: dato de la fuente (corrida del motor, libro de aforos, maestro). Negro: fórmula. Gris: columnas de "
                       "cálculo de las hojas de resumen.").font = FNOTE

    def R(n):
        return f"Detalle!${K[n]}$2:${K[n]}${n1}"

    meses = list(dict.fromkeys(d.mes))
    años = sorted(d.anio.unique())
    alcances = [(m, f'{R("Mes previsto")},"{m}"') for m in meses] + \
               [(f"Total {a}", f'{R("Año de cosecha")},{a}') for a in años]
    ambos = f'{R("Tiene los dos")},"Sí"'
    t = d[d.motor.notna() & d.aforo.notna()]
    con_dos = set(t.mes) | {f"Total {a}" for a in t.anio.unique()}
    alcances_dos = [(nom, cr) for nom, cr in alcances if nom in con_dos]  # sin aforos todavía, no hay qué comparar

    # ---- Resumen
    rs = wb.create_sheet("Resumen", 0)
    rs["A1"] = "Riopaila · Aforo de agronomía frente al motor · suertes que faltan por cosechar"
    rs["A1"].font = FTIT
    rs["A2"] = (f"Generado el {hoy.strftime('%d/%m/%Y')}. Sin TCH real todavía: se compara el aforo con el motor en las suertes "
                "que tienen los dos datos del ciclo en curso. TCH ponderado por área neta. Las columnas grises de la derecha dicen "
                f"cómo les fue a los dos contra el real en las cosechas cerradas desde el {desde_ref} ({n_ref} suertes).")
    rs["A2"].font, rs["A2"].alignment = FNOTE, Alignment(wrap_text=True)
    rs.merge_cells("A2:R2")
    rs.row_dimensions[2].height = 30
    rs.cell(4, 1, "Cobertura").font = FN
    for i, h in enumerate(["Mes previsto", "Suertes", "Área (ha)", "Con motor", "Con aforo del ciclo", "Con los dos",
                           "Toneladas del motor (todas las que tienen motor)"], 1):
        c = rs.cell(5, i, h)
        c.font, c.alignment, c.fill = FH, CEN, HEAD
    rs.row_dimensions[5].height = 40
    k = 6
    for nom, cr in alcances:
        rs.cell(k, 1, nom).font = FN if nom.startswith("Total") else FB
        vals = [f"=COUNTIFS({cr})", f"=SUMIFS({R('Área neta (ha)')},{cr})", f'=COUNTIFS({cr},{R("TCH motor")},">0")',
                f'=COUNTIFS({cr},{R("TCH aforo (Promedio)")},">0")', f"=COUNTIFS({cr},{ambos})",
                f"=SUMIFS({R('Ton motor (todas)')},{cr})"]
        for i, (x, fm) in enumerate(zip(vals, ["0", "#,##0.0", "0", "0", "0", "#,##0"]), 2):
            c = rs.cell(k, i, x)
            c.number_format, c.font = fm, FB
        k += 1
    fila = k + 1
    rs.cell(fila, 1, "Plantillas y socas (suertes con los dos datos)").font = FN
    heads = ["Mes previsto", "Grupo", "Corte desde", "Corte hasta", "Suertes", "Área (ha)", "TCH motor", "TCH aforo",
             "Aforo frente al motor", "Diferencia promedio por suerte", "Dentro de ±10 %", "Toneladas motor",
             "Toneladas aforo", "Suertes para revisar (> 15 %)", "Ref. sep-oct: aforo vs real", "Ref. sep-oct: motor vs real",
             "Ref. sep-oct: suertes"]
    for i, h in enumerate(heads, 1):
        c = rs.cell(fila + 1, i, h)
        c.font, c.alignment = FH, CEN
        c.fill = FREF if i >= 15 else HEAD
        if i >= 15:
            c.font = Font(name=F, size=10, bold=True, color="404040")
    rs.row_dimensions[fila + 1].height = 44
    k = fila + 2
    primera = k
    for nom, cr in alcances_dos:
        for g, lo, hi in BANDAS:
            tot = g == "Total"
            rs.cell(k, 1, nom).font = FN if tot else FB
            rs.cell(k, 2, g).font = FN if tot else FB
            rs.cell(k, 3, lo).font = rs.cell(k, 4, hi).font = FIN
            base = f'{ambos},{cr},{R("Corte")},">="&$C{k},{R("Corte")},"<="&$D{k}'
            s = lambda col: f"SUMIFS({R(col)},{base})"  # noqa: E731
            vals = [f"=COUNTIFS({base})", f"={s('Área neta (ha)')}", f'=IF(F{k}=0,"",{s("Motor × ha")}/F{k})',
                    f'=IF(F{k}=0,"",{s("Aforo × ha")}/F{k})', f'=IF(F{k}=0,"",H{k}/G{k}-1)',
                    f'=IF(E{k}=0,"",{s("|Aforo − motor| ÷ motor")}/E{k})', f"={s('Dentro de ±10 %')}",
                    f"={s('Motor × ha')}", f"={s('Aforo × ha')}", f'=COUNTIFS({base},{R("Revisar (dif. > 15 %)")},"Revisar")']
            for i, (x, fm) in enumerate(zip(vals, ["0", "#,##0.0", "0.0", "0.0", PCT, "0.0%", "0", "#,##0", "#,##0", "0"]), 5):
                c = rs.cell(k, i, x)
                c.number_format, c.font = fm, FB
            rv = ref.get(("banda", g.strip()))
            for i, val, fm in ((15, rv[0] if rv else None, PCT), (16, rv[1] if rv else None, PCT), (17, rv[2] if rv else None, "0")):
                c = rs.cell(k, i, val)
                c.number_format, c.font, c.fill = fm, FIN, FREF
            if tot:
                for col in range(1, 15):
                    rs.cell(k, col).fill = TOT
                    rs.cell(k, col).font = FN
            k += 1
        k += 1
    for i, w in enumerate([16, 22, 7, 7, 8, 9, 8, 8, 9, 10, 8, 10, 10, 10, 10, 10, 8], 1):
        rs.column_dimensions[get_column_letter(i)].width = w
    rs.freeze_panes = f"C{primera}"

    # ---- Variedad
    vs = wb.create_sheet("Variedad", 1)
    vs["A1"] = "Por variedad · suertes con los dos datos"
    vs["A1"].font = FTIT
    vs["A2"] = "La referencia de sep-oct aparece solo si la variedad tuvo 3 suertes o más con real, motor y aforo."
    vs["A2"].font = FNOTE
    vh = ["Alcance", "Variedad", "Suertes", "Área (ha)", "TCH motor", "TCH aforo", "Aforo frente al motor",
          "Diferencia promedio por suerte", "Dentro de ±10 %", "Suertes para revisar (> 15 %)",
          "Ref. sep-oct: aforo vs real", "Ref. sep-oct: motor vs real", "Ref. sep-oct: suertes"]
    for i, h in enumerate(vh, 1):
        c = vs.cell(4, i, h)
        c.font, c.alignment = FH, CEN
        c.fill = FREF if i >= 11 else HEAD
        if i >= 11:
            c.font = Font(name=F, size=10, bold=True, color="404040")
    vs.row_dimensions[4].height = 44
    k = 5
    for nom, cr in [(f"Total {a}", f'{R("Año de cosecha")},{a}') for a in años if f"Total {a}" in con_dos]:
        sub = t[t.anio == int(nom.split()[-1])]
        for var, _ in sorted(sub.variedad.value_counts().items(), key=lambda kv: (-kv[1], kv[0])):
            vs.cell(k, 1, nom).font, vs.cell(k, 2, var).font = FB, FB
            base = f'{ambos},{cr},{R("Variedad")},$B{k}'
            s = lambda col: f"SUMIFS({R(col)},{base})"  # noqa: E731
            vals = [f"=COUNTIFS({base})", f"={s('Área neta (ha)')}", f'=IF(D{k}=0,"",{s("Motor × ha")}/D{k})',
                    f'=IF(D{k}=0,"",{s("Aforo × ha")}/D{k})', f'=IF(D{k}=0,"",F{k}/E{k}-1)',
                    f'=IF(C{k}=0,"",{s("|Aforo − motor| ÷ motor")}/C{k})', f"={s('Dentro de ±10 %')}",
                    f'=COUNTIFS({base},{R("Revisar (dif. > 15 %)")},"Revisar")']
            for i, (x, fm) in enumerate(zip(vals, ["0", "#,##0.0", "0.0", "0.0", PCT, "0.0%", "0", "0"]), 3):
                c = vs.cell(k, i, x)
                c.number_format, c.font = fm, FB
            rv = ref.get(("var", var))
            for i, val, fm in ((11, rv[0] if rv else None, PCT), (12, rv[1] if rv else None, PCT), (13, rv[2] if rv else None, "0")):
                c = vs.cell(k, i, val)
                c.number_format, c.font, c.fill = fm, FIN, FREF
            k += 1
        k += 1
    for i, w in enumerate([12, 13, 8, 9, 8, 8, 9, 10, 8, 10, 10, 10, 8], 1):
        vs.column_dimensions[get_column_letter(i)].width = w
    vs.freeze_panes = "C5"

    # ---- Métodos
    me = wb.create_sheet("Métodos de aforo", 2)
    me["A1"] = "Métodos del aforo frente al motor · suertes con los dos datos"
    me["A1"].font = FTIT
    me["A2"] = "El TCH de agronomía («Promedio») es el promedio de los métodos Diámetro y Pesos."
    me["A2"].font = FNOTE
    for i, h in enumerate(["Alcance", "Método", "Suertes", "TCH del método", "TCH motor", "Frente al motor"], 1):
        c = me.cell(4, i, h)
        c.font, c.alignment, c.fill = FH, CEN, HEAD
    k = 5
    for nom, cr in alcances_dos:
        for mn, col in (("Diámetro", "Diámetro × ha"), ("Diámetro a 12 meses", "Diám. 12 m × ha"), ("Pesos", "Pesos × ha"),
                        ("Promedio (TCH de agronomía)", "Aforo × ha")):
            base = f"{ambos},{cr}"
            ha = f"SUMIFS({R('Área neta (ha)')},{base})"
            vals = [nom, mn, f"=COUNTIFS({base})", f'=IF(C{k}=0,"",SUMIFS({R(col)},{base})/{ha})',
                    f'=IF(C{k}=0,"",SUMIFS({R("Motor × ha")},{base})/{ha})', f'=IF(C{k}=0,"",D{k}/E{k}-1)']
            for i, (x, fm) in enumerate(zip(vals, [None, None, "0", "0.0", "0.0", PCT]), 1):
                c = me.cell(k, i, x)
                c.font = FN if mn.startswith("Promedio") else FB
                if fm:
                    c.number_format = fm
            k += 1
        k += 1
    for i, w in enumerate([16, 28, 9, 11, 10, 11], 1):
        me.column_dimensions[get_column_letter(i)].width = w

    # ---- Para revisar
    rv = wb.create_sheet("Para revisar", 3)
    rv["A1"] = "Suertes donde el aforo y el motor difieren más de 15 %, por toneladas en juego"
    rv["A1"].font = FTIT
    rv["A2"] = ("Foto al generar el libro (no se recalcula). «Toneladas en juego» = |aforo − motor| × área. La última columna dice "
                "cuál de los dos se acercó más al real en sep-oct en ese mismo grupo de corte.")
    rv["A2"].font = FNOTE
    t = t.assign(dif=t.aforo / t.motor - 1, enjuego=(t.aforo - t.motor).abs() * t.area)
    t = t[t.dif.abs() > 0.15].sort_values("enjuego", ascending=False)
    hh = ["Suerte", "Hacienda", "Técnico", "Variedad", "Corte", "Mes previsto", "Área (ha)", "TCH motor", "TCH aforo",
          "Aforo frente al motor", "Toneladas en juego", "Fecha del aforo", "En sep-oct, en su grupo de corte, se acercó más"]
    for i, h in enumerate(hh, 1):
        c = rv.cell(4, i, h)
        c.font, c.alignment, c.fill = FH, CEN, HEAD
    rv.row_dimensions[4].height = 32
    for k, x in enumerate(t.itertuples(index=False), start=5):
        bnom = "Plantilla (corte 1)" if x.corte == 1 else next(n.strip() for n, lo, hi in BANDAS[3:] if lo <= x.corte <= hi)
        rb = ref.get(("banda", bnom))
        quien = "" if not rb else ("el aforo" if abs(rb[0]) < abs(rb[1]) else "el motor") + \
            f" (aforo {rb[0]:+.0%}, motor {rb[1]:+.0%})"
        vals = [x.sec_ste, x.finca, x.tecnico, x.variedad, x.corte, x.mes, x.area, x.motor, x.aforo, x.dif, x.enjuego,
                x.f_eval, quien]
        for i, (v, fm) in enumerate(zip(vals, [None, None, None, None, "0", None, "0.00", "0.0", "0.0", PCT, "#,##0",
                                              "dd/mm/yyyy", None]), 1):
            c = rv.cell(k, i, v)
            c.font = FB
            if fm:
                c.number_format = fm
    for i, w in enumerate([10, 18, 20, 12, 6, 14, 8, 8, 8, 9, 10, 11, 34], 1):
        rv.column_dimensions[get_column_letter(i)].width = w
    rv.freeze_panes = "B5"

    # ---- Notas
    n = wb.create_sheet("Notas")
    n.column_dimensions["A"].width = 120
    sin_af = d[(d.motor.notna()) & d.aforo.isna()]
    notas = [
        ("Cómo se armó", True),
        (f"Suertes: maestro oficial ({fuente_maestro}), Riopaila, uso CAÑA, sin las que están en renovación, con próximo "
         f"corte en {hoy.year} (las que faltan por cosechar) o en {hoy.year + 1}. El mes es el de la fecha prevista del maestro.", False),
        (f"Motor: corrida del {json.load(open(rutas.TABLERO_DATA / 'meta.json', encoding='utf-8')).get('fecha_corrida')}, "
         "estimado del mismo corte que se va a cosechar (el ciclo en curso), con su factor de calibración.", False),
        ("Aforo: libro «Aforos con Diámetro» de Agronomía, columna «Promedio» (promedio de TCH Diámetro y TCH Pesos). Último "
         "aforo del ciclo en curso: mismo número de corte y evaluado después del último corte o la siembra.", False),
        (f"Los aforos se hacen hacia los 10 meses: {len(sin_af)} suertes con motor todavía no tienen aforo de su ciclo. Las de "
         f"{hoy.year + 1} irán entrando a medida que Agronomía las afore; basta volver a correr el programa.", False),
        (f"Referencia: cómo les fue al aforo y al motor contra el real, por grupo, en las cosechas cerradas desde el {desde_ref} "
         f"({n_ref} suertes con los tres datos; scripts/comparar_real_motor_aforo.py). Solo se muestra con 3 suertes o más.", False),
        ("", False),
        ("Cómo leer", True),
        ("Aforo frente al motor: TCH ponderado del aforo ÷ TCH ponderado del motor − 1. Positivo = el aforo espera más caña.", False),
        ("Diferencia promedio por suerte: promedio de |aforo − motor| ÷ motor, suerte por suerte.", False),
        ("Revisar: suertes con más de 15 % de diferencia entre los dos; la hoja «Para revisar» las ordena por toneladas en juego.", False),
        ("Generado con tch-motor/scripts/comparar_aforo_motor_pendientes.py.", False),
    ]
    for i, (txt, neg) in enumerate(notas, 1):
        c = n.cell(i, 1, txt)
        c.font, c.alignment = (FN if neg else FB), Alignment(wrap_text=True, vertical="top")
    for w_ in wb.worksheets:
        w_.sheet_view.showGridLines = w_.title == "Detalle"
    wb.active = 0
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--referencia-desde", default=f"{date.today().year}-09-01",
                    help="inicio de las cosechas cerradas con que se mide a aforo y motor contra el real")
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    hoy = date.today()
    salida = Path(a.salida) if a.salida else rutas.ESTIMADOS_DIR / f"Aforo_vs_motor_Riopaila_por_cosechar_{hoy.isoformat()}.xlsx"
    d, fuente_maestro = armar(hoy)
    ref, n_ref = referencia(a.referencia_desde, hoy.isoformat())
    t = d[d.motor.notna() & d.aforo.notna()]
    pond = lambda g, k: (g[k] * g.area).sum() / g.area.sum()  # noqa: E731
    for nom, g in list(t.groupby("mes", sort=False)) + [(f"Total", t)]:
        dif = (g.aforo - g.motor).abs() / g.motor
        print(f"{nom:16s} {len(g):>3} suertes · {g.area.sum():7.1f} ha · motor {pond(g, 'motor'):6.1f} · aforo {pond(g, 'aforo'):6.1f} "
              f"({pond(g, 'aforo') / pond(g, 'motor') - 1:+.1%}) · dif. por suerte {dif.mean():.1%} · revisar {int((dif > 0.15).sum())}")
    print("cobertura por año:", d.groupby("anio").apply(lambda g: f"{len(g)} suertes, {int(g.motor.notna().sum())} con motor, "
                                                         f"{int(g.aforo.notna().sum())} con aforo", include_groups=False).to_dict())
    escribir(d, ref, n_ref, a.referencia_desde, salida, hoy, fuente_maestro)
    print("guardado:", salida)


if __name__ == "__main__":
    main()
