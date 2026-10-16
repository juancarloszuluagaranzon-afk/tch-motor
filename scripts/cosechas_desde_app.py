# -*- coding: utf-8 -*-
"""Agrega a cosechas_recientes.xlsx las cosechas terminadas que el ZQM todavía no trae.

La corrida mensual (scripts/mensual.py) lee ese Excel como complemento del ZQM; sin las cosechas del mes,
el factor de calibración de la campaña no las ve, y el motor seguiría estimando como caña en pie una suerte que
ya se cortó si el maestro no ha registrado el corte.

  --mes AAAA-MM   cortes de Riopaila terminados ese mes en la app de Turnos (TCH de la app).
  --lista archivo una lista pegada, p. ej. la cosecha de Castilla, que no está en la app: suerte, TCH real y
                  fecha en que terminó (d/m/aaaa) por línea.

Antes de escribir guarda un respaldo del Excel en la misma carpeta (uno por día). No duplica: se salta la
suerte que ya tiene una fila con la misma fecha.

Uso:  .venv-gee\\Scripts\\python.exe scripts\\cosechas_desde_app.py --mes 2026-10 [--simular]
      .venv-gee\\Scripts\\python.exe scripts\\cosechas_desde_app.py --lista castilla_octubre.txt [--simular]
"""
import argparse
import calendar
import os
import re
import shutil
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl  # noqa: E402

from tchmotor import rutas, turnos  # noqa: E402


def leer_lista(ruta):
    """Lista pegada (p. ej. la cosecha de Castilla): una suerte por línea con suerte, TCH real y fecha en que
    terminó (d/m/aaaa), separados por tabulador, punto y coma o espacios. Las líneas sin esa forma (encabezados)
    se saltan."""
    out = []
    for linea in Path(ruta).read_text(encoding="utf-8-sig").splitlines():
        partes = [p for p in re.split(r"[\t;]+|\s{2,}|\s(?=\d)", linea.strip()) if p]
        if len(partes) < 3 or not re.fullmatch(r"\d{4}-\d{3}", partes[0]):
            continue
        try:
            tch = float(partes[1].replace(",", "."))
            fin = datetime.strptime(partes[2], "%d/%m/%Y").date()
        except ValueError:
            continue
        out.append({"sec_ste": partes[0], "tch_real": tch, "fecha_fin": fin.isoformat()})
    return out


def main():
    ap = argparse.ArgumentParser()
    fuente = ap.add_mutually_exclusive_group(required=True)
    fuente.add_argument("--mes", help="AAAA-MM: cortes de Riopaila terminados ese mes en la app de Turnos")
    fuente.add_argument("--lista", help="archivo con suerte, TCH real y fecha de terminación (d/m/aaaa) por línea")
    ap.add_argument("--simular", action="store_true", help="muestra lo que agregaría sin escribir")
    a = ap.parse_args()
    if a.lista:
        cortes = leer_lista(a.lista)
        print(f"lista {Path(a.lista).name}: {len(cortes)} cosechas leídas")
    else:
        y, m = map(int, a.mes.split("-"))
        desde, hasta = f"{a.mes}-01", f"{a.mes}-{calendar.monthrange(y, m)[1]:02d}"
        cortes = [c for c in turnos.cortes_terminados(desde, hasta) if c["tch_real"] is not None]
        print(f"app de Turnos: {len(cortes)} cortes de Riopaila terminados en {a.mes} con TCH")

    arch = rutas.COSECHAS_RECIENTES
    wb = openpyxl.load_workbook(arch)
    ws = wb["Cosechas"]
    ya = {(str(r[0]).strip(), r[2].date() if hasattr(r[2], "date") else r[2])
          for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}
    nuevos = [c for c in sorted(cortes, key=lambda x: (x["fecha_fin"], x["sec_ste"]))
              if (c["sec_ste"], date.fromisoformat(c["fecha_fin"])) not in ya]
    for c in nuevos:
        print(f"  {c['sec_ste']}  TCH {c['tch_real']:.2f}  termina {c['fecha_fin']}")
    if a.simular or not nuevos:
        print("nada escrito" + (" (simulación)" if a.simular else ": ya estaban todas"))
        return
    resp = arch.with_name(f"{arch.stem}_respaldo_{date.today().isoformat()}{arch.suffix}")
    if not resp.exists():
        shutil.copy2(arch, resp)
        print("respaldo:", resp.name)
    for c in nuevos:
        k = ws.max_row + 1
        ws.cell(k, 1, c["sec_ste"])
        ws.cell(k, 2, round(c["tch_real"], 2))
        ws.cell(k, 3, datetime.fromisoformat(c["fecha_fin"])).number_format = "d/mm/yyyy"
    wb.save(arch)
    print(f"agregadas {len(nuevos)} cosechas · {arch}")


if __name__ == "__main__":
    main()
