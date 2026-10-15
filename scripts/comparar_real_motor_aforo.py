# -*- coding: utf-8 -*-
"""TCH real (app de Turnos) frente al motor y al aforo de Agronomía, suerte por suerte. Riopaila. Excel con fórmulas.

Real   = cortes terminados del módulo de Cosecha (tchmotor.turnos: área efectiva y caña total del corte);
         el mes es el de la fecha de cierre.
Motor  = el estimado vigente ANTES del corte y del MISMO número de corte: la corrida más reciente del historial
         del tablero anterior a la fecha de cierre; si no hay, la estimación del motor sin conocer el real
         (reales.json) por el factor de calibración de las cosechas de los meses anteriores.
Aforo  = columna «Promedio» (promedio de TCH Diámetro y TCH Pesos) del libro «Aforos con Diámetro» de Agronomía:
         el último aforo del mismo corte antes del cierre (máximo 270 días).
Las hojas de resumen comparan solo las suertes que tienen los tres datos.

Uso:  .venv-gee\\Scripts\\python.exe scripts\\comparar_real_motor_aforo.py --desde 2026-09-01 --hasta 2026-10-31 [--salida x.xlsx]
"""
import argparse
import json
import os
import shutil
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from tchmotor import rutas, turnos  # noqa: E402

PLANTA = "riopaila"
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre",
         "Noviembre", "Diciembre"]
TODOS = "Todos los meses"
RECAL_MIN = 15


# ======================================================================== datos
def motor_por_suerte(r):
    hist = {}
    for f in sorted((rutas.TABLERO_DATA / "historial").glob("estimados_*.json")):
        e = pd.DataFrame(json.load(open(f, encoding="utf-8")))
        hist[f.stem.split("_", 1)[1]] = e[e.planta == PLANTA].set_index("sec_ste")
    rj = pd.DataFrame(json.load(open(rutas.TABLERO_DATA / "reales.json", encoding="utf-8")))
    rj = rj[rj.planta == PLANTA]

    def factor_antes(anio, mes):
        p = rj[(rj.anio == anio) & (rj.mes < mes)]
        if len(p) < RECAL_MIN:
            return 1.0
        return float((p.tch_real * p.area_ha).sum() / (p.tch_modelo * p.area_ha).sum())

    out = []
    for x in r.itertuples():
        res = (np.nan, np.nan, None)
        for fecha in sorted(hist, reverse=True):
            if fecha >= x.fecha_fin:
                continue
            e = hist[fecha]
            if x.sec_ste in e.index and e.loc[x.sec_ste, "soca"] == x.corte:
                est = float(e.loc[x.sec_ste, "tch_estimado"])
                res = (est, est / float(e.loc[x.sec_ste, "factor_recal"]), f"Corrida {fecha}")
                break
        if res[2] is None:
            fin = date.fromisoformat(x.fecha_fin)
            q = rj[(rj.sec_ste == x.sec_ste) & (rj.corte == x.corte) & (rj.anio == fin.year)]
            if len(q):
                crudo, k = float(q.tch_modelo.iloc[-1]), factor_antes(fin.year, fin.month)
                res = (crudo * k, crudo, f"Estimación sin el real × factor {k:.4f}")
        out.append(res)
    return out


def aforos_por_suerte(r):
    src = rutas.AFOROS_DIAMETRO
    if not src.exists():
        raise FileNotFoundError(f"No encuentro el libro de aforos en {src} (ajusta TCH_AFOROS_DIAMETRO)")
    with tempfile.TemporaryDirectory() as tmp:  # OneDrive bloquea la lectura directa: se lee una copia
        copia = Path(tmp) / "aforos.xlsx"
        shutil.copy2(src, copia)
        a = pd.read_excel(copia, "Aforos", header=1)
    a = a.rename(columns={a.columns[0]: "sec", a.columns[3]: "f_eval", a.columns[5]: "corte_af",
                          a.columns[13]: "af_diam", a.columns[14]: "af_diam12", a.columns[15]: "af_pesos",
                          a.columns[16]: "aforo"})
    a["sec"] = a.sec.astype(str).str.strip()
    a["f_eval"] = pd.to_datetime(a.f_eval, errors="coerce")
    c = a.merge(r[["sec_ste", "fin", "corte"]], left_on="sec", right_on="sec_ste")
    c = c[(c.f_eval < c.fin) & ((c.fin - c.f_eval).dt.days <= 270) & (c.corte_af == c.corte)]
    c = c.sort_values("f_eval").groupby("sec_ste").tail(1)
    c["dias_af"] = (c.fin - c.f_eval).dt.days
    return c[["sec_ste", "aforo", "af_diam", "af_diam12", "af_pesos", "f_eval", "dias_af"]]


def armar(desde, hasta):
    r = pd.DataFrame(turnos.cortes_terminados(desde, hasta))
    r = r[r.area_ef.notna()].copy()
    r["corte"] = r.corte.astype(int)
    r["fin"] = pd.to_datetime(r.fecha_fin)
    r["mes"] = [f"{MESES[d.month - 1]} {d.year}" for d in r.fin]
    r[["motor", "motor_crudo", "motor_fuente"]] = pd.DataFrame(motor_por_suerte(r), index=r.index)
    r = r.merge(aforos_por_suerte(r), on="sec_ste", how="left")
    r["banda"] = pd.cut(r.corte, [0, 1, 3, 6, 10, 99], labels=["Plantilla", "Soca 2-3", "Soca 4-6", "Soca 7-10", "Soca 11+"]).astype(str)
    return r.sort_values(["fecha_fin", "sec_ste"]).reset_index(drop=True)


def resumen_consola(r):
    t = r[r.motor.notna() & r.aforo.notna()]
    pond = lambda g, k: (g[k] * g.area_ef).sum() / g.area_ef.sum()  # noqa: E731
    for nom, g in list(t.groupby("mes", sort=False)) + [(TODOS, t)]:
        em, ea = (g.motor - g.tch_real) / g.tch_real, (g.aforo - g.tch_real) / g.tch_real
        print(f"{nom:16s} {len(g):>3} suertes · real {pond(g, 'tch_real'):6.1f} · motor {pond(g, 'motor') / pond(g, 'tch_real') - 1:+6.1%} "
              f"(error {em.abs().mean():5.1%}) · aforo {pond(g, 'aforo') / pond(g, 'tch_real') - 1:+6.1%} (error {ea.abs().mean():5.1%}) · "
              f"más cerca motor {int((em.abs() < ea.abs()).sum())}-{int((em.abs() >= ea.abs()).sum())}")


# ======================================================================== Excel
F = "Arial"
f_base, f_bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
f_head, f_title = Font(name=F, size=10, bold=True, color="FFFFFF"), Font(name=F, size=14, bold=True)
f_note, f_input = Font(name=F, size=9, italic=True, color="595959"), Font(name=F, size=10, color="0000FF")
fill_head, fill_calc, fill_tot = (PatternFill("solid", fgColor=c) for c in ("2F3A33", "EFEFEF", "F4F4F0"))
fill_real, fill_motor, fill_aforo = (PatternFill("solid", fgColor=c) for c in ("DCE9F8", "FBE3D6", "D5F2E7"))
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
PCT = "+0.0%;-0.0%;0.0%"


def escribir(r, salida, desde, hasta):
    wb = Workbook()
    d = wb.active
    d.title = "Detalle"
    cols = [("Suerte", 10, None), ("Hacienda", 18, None), ("Zona", 6, "0"), ("Variedad", 12, None), ("Corte", 6, "0"),
            ("Tipo", 10, None), ("Grupo de corte", 11, None), ("Mes de cierre", 14, None), ("Fecha de cierre", 11, "dd/mm/yyyy"),
            ("Área efectiva (ha)", 10, "0.00"), ("Toneladas", 10, "#,##0.0"), ("TCH real", 9, "0.0"),
            ("TCH motor", 9, "0.0"), ("Fuente del motor", 26, None), ("TCH aforo (Promedio)", 10, "0.0"),
            ("Aforo Diámetro", 9, "0.0"), ("Aforo Diámetro 12 m", 9, "0.0"), ("Aforo Pesos", 9, "0.0"),
            ("Fecha del aforo", 11, "dd/mm/yyyy"), ("Días aforo → cierre", 9, "0"),
            ("Tiene los tres datos", 9, None), ("Error motor", 9, PCT), ("Error aforo", 9, PCT), ("Más cerca del real", 10, None),
            ("Real × ha", 10, "#,##0.0"), ("Motor × ha", 10, "#,##0.0"), ("Aforo × ha", 10, "#,##0.0"),
            ("Diámetro × ha", 10, "#,##0.0"), ("Diám. 12 m × ha", 10, "#,##0.0"), ("Pesos × ha", 10, "#,##0.0"),
            ("|Error| motor", 9, "0.0%"), ("|Error| aforo", 9, "0.0%"), ("|Error| diámetro", 9, "0.0%"),
            ("|Error| diám. 12 m", 9, "0.0%"), ("|Error| pesos", 9, "0.0%"),
            ("Motor ±10 %", 8, "0"), ("Aforo ±10 %", 8, "0"), ("Diámetro ±10 %", 8, "0"), ("Diám. 12 m ±10 %", 8, "0"),
            ("Pesos ±10 %", 8, "0"), ("Gana motor", 8, "0")]
    C = {n: get_column_letter(i + 1) for i, (n, _, _) in enumerate(cols)}
    n0, n1 = 2, len(r) + 1
    for i, (n, w, _) in enumerate(cols, 1):
        c = d.cell(1, i, n)
        c.alignment = center
        c.fill, c.font = (fill_calc, Font(name=F, size=9, bold=True, color="404040")) if i >= 25 else (fill_head, f_head)
        d.column_dimensions[get_column_letter(i)].width = w
    d.row_dimensions[1].height = 42
    nan = lambda v: None if (v is None or (isinstance(v, float) and v != v)) else v  # noqa: E731
    for k, x in enumerate(r.itertuples(index=False), start=2):
        vals = [x.sec_ste, x.hacienda, x.zona, x.variedad, x.corte, "Plantilla" if x.corte == 1 else "Soca",
                "Plantilla" if x.corte == 1 else x.banda, x.mes, x.fin.to_pydatetime(), x.area_ef, x.ton, None,
                nan(x.motor), x.motor_fuente, nan(x.aforo), nan(x.af_diam), nan(x.af_diam12), nan(x.af_pesos),
                x.f_eval.to_pydatetime() if isinstance(x.f_eval, pd.Timestamp) else None, nan(x.dias_af)]
        for j, v in enumerate(vals, 1):
            d.cell(k, j, nan(v))
        Lr, J, K, M, O = (f"{C[n]}{k}" for n in ("TCH real", "Área efectiva (ha)", "Toneladas", "TCH motor", "TCH aforo (Promedio)"))
        P, Q, R = (f"{C[n]}{k}" for n in ("Aforo Diámetro", "Aforo Diámetro 12 m", "Aforo Pesos"))
        U, X = f"{C['Tiene los tres datos']}{k}", f"{C['Más cerca del real']}{k}"
        f = {"TCH real": f"={K}/{J}",
             "Tiene los tres datos": f'=IF(AND(ISNUMBER({M}),ISNUMBER({O})),"Sí","No")',
             "Error motor": f'=IF(ISNUMBER({M}),{M}/{Lr}-1,"")', "Error aforo": f'=IF(ISNUMBER({O}),{O}/{Lr}-1,"")',
             "Más cerca del real": f'=IF({U}="Sí",IF(ABS({M}-{Lr})<ABS({O}-{Lr}),"Motor","Aforo"),"")',
             "Real × ha": f"={Lr}*{J}", "Gana motor": f'=IF({X}="Motor",1,0)'}
        for nom, ref in (("Motor", M), ("Aforo", O), ("Diámetro", P), ("Diám. 12 m", Q), ("Pesos", R)):
            f[f"{nom} × ha"] = f"=IF(ISNUMBER({ref}),{ref}*{J},0)"
            f[f"|Error| {nom.lower()}"] = f"=IF(ISNUMBER({ref}),ABS({ref}-{Lr})/{Lr},0)"
            f[f"{nom} ±10 %"] = f"=IF(ISNUMBER({ref}),IF(ABS({ref}-{Lr})/{Lr}<=0.1,1,0),0)"
        for nom, frm in f.items():
            d[f"{C[nom]}{k}"] = frm
        for j, (_, _, fmt) in enumerate(cols, 1):
            cell = d.cell(k, j)
            cell.font = f_base if j <= 24 else Font(name=F, size=9, color="595959")
            if fmt:
                cell.number_format = fmt
            if j in (10, 11, 13, 15, 16, 17, 18) and cell.value is not None:
                cell.font = f_input
    d.freeze_panes = "B2"
    d.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n1}"
    d.cell(n1 + 2, 1, "Azul: dato de la fuente (app de Cosecha, corridas del motor, libro de aforos). Negro: fórmula. "
                      "Gris: columnas de cálculo de las hojas de resumen.").font = f_note

    def rng(nom):
        return f"Detalle!${C[nom]}${n0}:${C[nom]}${n1}"

    meses = list(dict.fromkeys(r.mes)) + [TODOS]
    crit_mes = lambda cell: f'IF({cell}="{TODOS}","*",{cell})'  # noqa: E731
    METR = ["Suertes", "Área (ha)", "TCH real", "TCH motor", "Desvío motor", "Error por suerte motor", "Motor dentro de ±10 %",
            "TCH aforo", "Desvío aforo", "Error por suerte aforo", "Aforo dentro de ±10 %", "Más cerca: motor",
            "Más cerca: aforo", "Aforo frente al motor"]
    FMT = ["0", "#,##0.0", "0.0", "0.0", PCT, "0.0%", "0", "0.0", PCT, "0.0%", "0", "0", "0", PCT]

    def metricas(ws, k, c0, criterios):
        L = lambda i: f"{get_column_letter(c0 + i)}{k}"  # noqa: E731
        base = f'{rng("Tiene los tres datos")},"Sí",{criterios}'
        s = lambda col: f"SUMIFS({rng(col)},{base})"  # noqa: E731
        frm = [f"=COUNTIFS({base})", f"={s('Área efectiva (ha)')}", f'=IF({L(1)}=0,"",{s("Real × ha")}/{L(1)})',
               f'=IF({L(1)}=0,"",{s("Motor × ha")}/{L(1)})', f'=IF({L(1)}=0,"",{L(3)}/{L(2)}-1)',
               f'=IF({L(0)}=0,"",{s("|Error| motor")}/{L(0)})', f"={s('Motor ±10 %')}",
               f'=IF({L(1)}=0,"",{s("Aforo × ha")}/{L(1)})', f'=IF({L(1)}=0,"",{L(7)}/{L(2)}-1)',
               f'=IF({L(0)}=0,"",{s("|Error| aforo")}/{L(0)})', f"={s('Aforo ±10 %')}", f"={s('Gana motor')}",
               f"={L(0)}-{L(11)}", f'=IF({L(1)}=0,"",{L(7)}/{L(3)}-1)']
        for i, (x, fmt) in enumerate(zip(frm, FMT)):
            c = ws.cell(k, c0 + i, x)
            c.number_format, c.font = fmt, f_base

    def encabezado(ws, fila, campos, c0):
        for txt, a, b, fill in [("Real", c0, c0 + 2, fill_real), ("Motor", c0 + 3, c0 + 6, fill_motor),
                                ("Aforo agronomía", c0 + 7, c0 + 10, fill_aforo), ("¿Quién quedó más cerca?", c0 + 11, c0 + 13, fill_tot)]:
            ws.merge_cells(start_row=fila, start_column=a, end_row=fila, end_column=b)
            c = ws.cell(fila, a, txt)
            c.font, c.alignment = f_bold, center
            for col in range(a, b + 1):
                ws.cell(fila, col).fill = fill
        for i, n in enumerate(campos, 1):
            c = ws.cell(fila + 1, i, n)
            c.font, c.alignment, c.fill = f_head, center, fill_head
        ws.row_dimensions[fila + 1].height = 40

    # ---- Resumen
    rs = wb.create_sheet("Resumen", 0)
    rs["A1"] = f"Riopaila · TCH real frente al motor y al aforo de agronomía · cosechas terminadas del {desde} al {hasta}"
    rs["A1"].font = f_title
    rs["A2"] = ("Las comparaciones usan solo las suertes con real, motor y aforo. TCH ponderado por área efectiva. "
                "Desvío = estimado ÷ real − 1. Error por suerte = promedio de |estimado − real| ÷ real.")
    rs["A2"].font = f_note
    rs.cell(4, 1, "Cobertura").font = f_bold
    camp = ["Mes", "Terminadas", "Área (ha)", "TCH real (todas)", "Con motor", "Con aforo", "Con los tres datos",
            "Desvío motor (todas las que tienen motor)", "Error por suerte motor (todas)",
            "Desvío aforo (todas las que tienen aforo)", "Error por suerte aforo (todas)"]
    for i, n in enumerate(camp, 1):
        c = rs.cell(5, i, n)
        c.font, c.alignment, c.fill = f_head, center, fill_head
    rs.row_dimensions[5].height = 52
    for j, mes in enumerate(meses):
        k = 6 + j
        rs.cell(k, 1, mes).font = f_bold
        H = f'{rng("Mes de cierre")},{crit_mes(f"$A{k}")}'
        mot, afo = f'{rng("TCH motor")},">0"', f'{rng("TCH aforo (Promedio)")},">0"'
        frm = [f"=COUNTIFS({H})", f"=SUMIFS({rng('Área efectiva (ha)')},{H})", f"=SUMIFS({rng('Real × ha')},{H})/C{k}",
               f"=COUNTIFS({H},{mot})", f"=COUNTIFS({H},{afo})", f'=COUNTIFS({H},{rng("Tiene los tres datos")},"Sí")',
               f'=IF(E{k}=0,"",SUMIFS({rng("Motor × ha")},{H},{mot})/SUMIFS({rng("Real × ha")},{H},{mot})-1)',
               f'=IF(E{k}=0,"",SUMIFS({rng("|Error| motor")},{H},{mot})/E{k})',
               f'=IF(F{k}=0,"",SUMIFS({rng("Aforo × ha")},{H},{afo})/SUMIFS({rng("Real × ha")},{H},{afo})-1)',
               f'=IF(F{k}=0,"",SUMIFS({rng("|Error| aforo")},{H},{afo})/F{k})']
        for i, (x, fm) in enumerate(zip(frm, ["0", "#,##0.0", "0.0", "0", "0", "0", PCT, "0.0%", PCT, "0.0%"]), 2):
            c = rs.cell(k, i, x)
            c.number_format, c.font = fm, f_base
    fila = 6 + len(meses) + 1
    rs.cell(fila, 1, "Plantillas y socas (suertes con los tres datos)").font = f_bold
    GR = [("Total", 1, 99), ("Plantilla (corte 1)", 1, 1), ("Socas (corte 2 o más)", 2, 99), ("   Soca 2-3", 2, 3),
          ("   Soca 4-6", 4, 6), ("   Soca 7-10", 7, 10), ("   Soca 11 o más", 11, 99)]
    encabezado(rs, fila + 1, ["Mes", "Grupo", "Corte desde", "Corte hasta"] + METR, 5)
    k = fila + 3
    primera = k
    for mes in meses:
        for g, lo, hi in GR:
            tot = g == "Total"
            rs.cell(k, 1, mes).font = f_bold if tot else f_base
            rs.cell(k, 2, g).font = f_bold if tot else f_base
            rs.cell(k, 3, lo).font = rs.cell(k, 4, hi).font = f_input
            metricas(rs, k, 5, f'{rng("Mes de cierre")},{crit_mes(f"$A{k}")},{rng("Corte")},">="&$C{k},{rng("Corte")},"<="&$D{k}')
            if tot:
                for col in range(1, 19):
                    rs.cell(k, col).fill = fill_tot
                    rs.cell(k, col).font = f_bold
            k += 1
        k += 1
    for i, w in enumerate([16, 22, 7, 7, 8, 9, 8, 8, 9, 9, 9, 8, 9, 9, 9, 8, 8, 9], 1):
        rs.column_dimensions[get_column_letter(i)].width = w
    rs.freeze_panes = f"C{primera}"

    # ---- Variedad
    v = wb.create_sheet("Variedad", 1)
    v["A1"] = "Por variedad · suertes con los tres datos"
    v["A1"].font = f_title
    v["A2"] = "Ordenadas por número de suertes en cada bloque. Con menos de 3 suertes la cifra es poco confiable."
    v["A2"].font = f_note
    encabezado(v, 3, ["Mes", "Variedad"] + METR, 3)
    tri = r[r.motor.notna() & r.aforo.notna()]
    k = 5
    for mes in meses:
        sub = tri if mes == TODOS else tri[tri.mes == mes]
        for var, _ in sorted(sub.variedad.value_counts().items(), key=lambda kv: (-kv[1], kv[0])):
            v.cell(k, 1, mes).font, v.cell(k, 2, var).font = f_base, f_base
            metricas(v, k, 3, f'{rng("Mes de cierre")},{crit_mes(f"$A{k}")},{rng("Variedad")},$B{k}')
            k += 1
        k += 1
    for i, w in enumerate([16, 13, 8, 9, 8, 8, 9, 9, 9, 8, 9, 9, 9, 8, 8, 9], 1):
        v.column_dimensions[get_column_letter(i)].width = w
    v.freeze_panes = "C5"

    # ---- Métodos
    m = wb.create_sheet("Métodos de aforo", 2)
    m["A1"] = "Métodos del aforo frente al real · suertes con los tres datos"
    m["A1"].font = f_title
    m["A2"] = "El TCH de agronomía («Promedio») es el promedio de los métodos Diámetro y Pesos."
    m["A2"].font = f_note
    for i, n in enumerate(["Mes", "Método", "Suertes", "TCH del método", "TCH real", "Desvío", "Error por suerte",
                           "Dentro de ±10 %"], 1):
        c = m.cell(4, i, n)
        c.font, c.alignment, c.fill = f_head, center, fill_head
    MET = [("Diámetro", "Diámetro"), ("Diámetro a 12 meses", "Diám. 12 m"), ("Pesos", "Pesos"),
           ("Promedio (TCH de agronomía)", "Aforo"), ("Motor", "Motor")]
    k = 5
    for mes in meses:
        for nom, cl in MET:
            base = f'{rng("Tiene los tres datos")},"Sí",{rng("Mes de cierre")},{crit_mes(f"$A{k}")}'
            ha = f"SUMIFS({rng('Área efectiva (ha)')},{base})"
            vals = [mes, nom, f"=COUNTIFS({base})", f'=IF(C{k}=0,"",SUMIFS({rng(cl + " × ha")},{base})/{ha})',
                    f'=IF(C{k}=0,"",SUMIFS({rng("Real × ha")},{base})/{ha})', f'=IF(C{k}=0,"",D{k}/E{k}-1)',
                    f'=IF(C{k}=0,"",SUMIFS({rng("|Error| " + cl.lower())},{base})/C{k})',
                    f"=SUMIFS({rng(cl + ' ±10 %')},{base})"]
            for i, (x, fm) in enumerate(zip(vals, [None, None, "0", "0.0", "0.0", PCT, "0.0%", "0"]), 1):
                c = m.cell(k, i, x)
                c.font = f_bold if nom.startswith(("Promedio", "Motor")) else f_base
                if fm:
                    c.number_format = fm
            k += 1
        k += 1
    for i, w in enumerate([16, 28, 9, 11, 9, 9, 10, 10], 1):
        m.column_dimensions[get_column_letter(i)].width = w

    # ---- Notas
    n = wb.create_sheet("Notas")
    n.column_dimensions["A"].width = 120
    sin_motor = r[r.motor.isna()].sec_ste.tolist()
    sin_aforo = r[r.aforo.isna()].sec_ste.tolist()
    notas = [("Cómo se armó", True),
             ("Real: módulo de Cosecha de la app de Turnos. Toneladas del corte completo (vagones, caña de meses anteriores y caña "
              "molida en otro ingenio) sobre el área efectiva (área del maestro menos la descontada). El mes es el de la fecha de cierre.", False),
             ("Motor: el estimado vigente antes del corte y del mismo número de corte. La corrida más reciente del historial del "
              "tablero anterior al cierre; si no hay ninguna, la estimación del motor sin conocer el real (reales.json) por el factor "
              "de calibración de las cosechas de los meses anteriores. La columna «Fuente del motor» dice cuál se usó.", False),
             ("Aforo: libro «Aforos con Diámetro» de Agronomía, columna «Promedio» (promedio de TCH Diámetro y TCH Pesos). Último "
              "aforo del mismo corte antes del cierre, como máximo 270 días antes.", False),
             (f"Sin motor ({len(sin_motor)}): {', '.join(sin_motor) or 'ninguna'}. Sin aforo del ciclo ({len(sin_aforo)}): "
              f"{', '.join(sin_aforo) or 'ninguna'}.", False),
             ("", False), ("Cómo leer las columnas", True),
             ("Desvío: TCH ponderado del estimado ÷ TCH ponderado real − 1. Positivo = el estimado se pasó.", False),
             ("Error por suerte: promedio de |estimado − real| ÷ real, suerte por suerte (MAPE).", False),
             ("Dentro de ±10 %: suertes cuyo estimado quedó a 10 % o menos de su real.", False),
             ("Más cerca: en cuántas suertes el motor o el aforo quedó más cerca del real (el empate cuenta para el aforo).", False),
             ("Las hojas de resumen son fórmulas sobre el Detalle: si se corrige un dato en azul, todo se recalcula.", False),
             ("Generado con tch-motor/scripts/comparar_real_motor_aforo.py.", False)]
    for i, (txt, neg) in enumerate(notas, 1):
        c = n.cell(i, 1, txt)
        c.font, c.alignment = (f_bold if neg else f_base), Alignment(wrap_text=True, vertical="top")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title == "Detalle"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--desde", required=True, help="AAAA-MM-DD (fecha de cierre)")
    ap.add_argument("--hasta", required=True, help="AAAA-MM-DD (fecha de cierre)")
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    salida = Path(a.salida) if a.salida else rutas.ESTIMADOS_DIR / f"Motor_aforo_vs_real_Riopaila_{a.desde}_{a.hasta}.xlsx"
    r = armar(a.desde, a.hasta)
    print(f"{len(r)} cortes terminados · con motor {int(r.motor.notna().sum())} · con aforo {int(r.aforo.notna().sum())} · "
          f"motor por fuente: {r.motor_fuente.value_counts().to_dict()}")
    resumen_consola(r)
    escribir(r, salida, a.desde, a.hasta)
    print("guardado:", salida)


if __name__ == "__main__":
    main()
