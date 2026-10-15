"""Exporta a Excel TODOS los estimados del modelo, todos los años.

- 2019–2024: estimado out-of-sample (RF por GroupKFold-año, nunca ve el año que
  predice) sobre los 504 eventos validados, con TCH real (etiqueta) y error.
- 2026: estimados de producción (tabla tch_estimado del Supabase de Rio Map, 839
  suertes) con TCH real donde ya cerró el corte.

Salida: C:\\Users\\Agr349\\Documents\\Estimados_TCH_modelo.xlsx
"""
import io, os, sys, json, numpy as np, pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sklearn.model_selection import GroupKFold

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
sys.path.insert(0, r"C:\Users\Agr349\tch-motor\scripts")
from tchmotor import (cargar_var_map, AforoSintetico, build_features, MotorTCH,
                      map_variedad, Supabase)
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

DS = r"C:\Users\Agr349\tch-motor\data\dataset_entrenamiento.jsonl"
AFORO = r"C:\Users\Agr349\Documents\Base de aforos completa 2026.xlsx"
OUT = r"C:\Users\Agr349\Documents\Estimados_TCH_modelo.xlsx"


def num(x):
    try:
        return float(str(x).replace(",", "."))
    except Exception:
        return np.nan


def registros_aforo():
    wb = openpyxl.load_workbook(AFORO, read_only=True, data_only=True)
    out = []
    for name in ["Aforos 2024-2026", "GENERAL DE LA SUERTE", "BASE 2022-2023", "AFOROS 2022-2023"]:
        if name not in wb.sheetnames:
            continue
        rows = list(wb[name].iter_rows(values_only=True, max_col=34))
        hi = next((i for i, r in enumerate(rows[:6]) if any(str(c).strip().upper() == "VARIEDAD" for c in r if c)), None)
        if hi is None:
            continue
        H = {str(c).strip().lower(): j for j, c in enumerate(rows[hi]) if c}
        fnd = lambda *s: next((j for k, j in H.items() if all(x in k for x in s)), None)
        cv, cc, cpob, ctch = fnd("variedad"), fnd("corte"), fnd("tallos", "mt"), fnd("tch", "calcul")
        cpe, cdi = fnd("total", "peso"), fnd("total", "diam")
        if cv is None or cpob is None:
            continue
        for r in rows[hi + 1:]:
            var = str(r[cv]).strip() if r[cv] else None
            pob = num(r[cpob])
            if not var or not (var.startswith("CC") or var == "VARIAS") or not (3 < pob < 25):
                continue
            out.append(dict(variedad=var, soca=num(r[cc]) if cc is not None else None,
                            pob=pob, peso=num(r[cpe]) if cpe else None,
                            dia=num(r[cdi]) if cdi else None, tch=num(r[ctch]) if ctch else None))
    wb.close()
    return out


def _env(path, name):
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if line.startswith(name + "="):
            return line.split("=", 1)[1].strip().strip('"').strip("'")
    return None


# ---- 2019-2024: estimado out-of-sample ----
vm = cargar_var_map()
sint = AforoSintetico(registros_aforo(), min_n=4)
feats, meta = [], []
for line in open(DS, encoding="utf-8"):
    r = json.loads(line)
    if not r["curves"].get("NDVI"):
        continue
    f = build_features(r, sint, vm)
    if f.get("ndvi_n", 0) < 3:
        continue
    feats.append(f); meta.append(r)
X = pd.DataFrame(feats)
y = np.array([m["tch"] for m in meta])
g = np.array([int(m["anio"]) for m in meta])
pred = np.zeros(len(y)); std = np.zeros(len(y))
for tr, te in GroupKFold(5).split(X, y, g):
    m = MotorTCH(n_estimators=600).entrenar(X.iloc[tr], y[tr])
    mu, sd = m.predecir_con_banda(X.iloc[te])
    pred[te] = mu; std[te] = sd

filas = []
for i, r in enumerate(meta):
    est = round(float(pred[i]), 1)
    real = round(float(r["tch"]), 1)
    filas.append({
        "anio": int(r["anio"]), "sec_ste": r["sec_ste"], "planta": r["planta"],
        "variedad": map_variedad(r.get("variedad"), vm), "soca": r.get("corte"),
        "edad_meses": round(float(r["edad"]), 1) if r.get("edad") is not None else None,
        "area_ha": round(float(r["ha"]), 2), "tch_estimado": est,
        "tch_std": round(float(std[i]), 1), "tch_real": real,
        "error_tha": round(est - real, 1), "error_pct": round((est - real) / real * 100, 1) if real else None,
        "modo": "motor", "fuente": "validación (fuera de muestra)",
    })
print(f"estimados validación 2019-2024: {len(filas)}")

# ---- 2026: producción (Supabase) ----
try:
    url = _env(r"C:\Users\Agr349\agrocontrol-campo\.env.local", "NEXT_PUBLIC_SUPABASE_URL")
    anon = _env(r"C:\Users\Agr349\agrocontrol-campo\.env.local", "NEXT_PUBLIC_SUPABASE_ANON_KEY")
    sb = Supabase(url=url, key=anon)
    est26 = sb.leer("tch_estimado", select="*")
    reales = {r["sec_ste"]: r for r in sb.leer("tch_real")}
    for r in est26:
        real = reales.get(r["sec_ste"], {}).get("tch_real")
        e = r["tch_estimado"]
        filas.append({
            "anio": 2026, "sec_ste": r["sec_ste"], "planta": r["planta"],
            "variedad": r.get("variedad"), "soca": r.get("soca"),
            "edad_meses": r.get("edad_meses"), "area_ha": r.get("area_ha"),
            "tch_estimado": e, "tch_std": r.get("tch_std"), "tch_real": real,
            "error_tha": round(e - real, 1) if real is not None else None,
            "error_pct": round((e - real) / real * 100, 1) if real else None,
            "modo": r.get("modo"), "fuente": "producción",
        })
    print(f"estimados producción 2026: {len(est26)} (con real: {sum(1 for r in est26 if r['sec_ste'] in reales)})")
except Exception as ex:
    print("aviso: no pude leer 2026 de Supabase:", str(ex)[:120])

df = pd.DataFrame(filas).sort_values(["anio", "planta", "sec_ste"]).reset_index(drop=True)

# ---- resumen por año (ponderado por área) ----
res = []
for a, gdf in df.groupby("anio"):
    w = gdf["area_ha"].fillna(0)
    est_agg = (gdf["tch_estimado"] * w).sum() / w.sum() if w.sum() else np.nan
    con_real = gdf.dropna(subset=["tch_real"])
    wr = con_real["area_ha"].fillna(0)
    real_agg = (con_real["tch_real"] * wr).sum() / wr.sum() if wr.sum() else np.nan
    res.append({"anio": a, "n_suertes": len(gdf), "area_ha": round(w.sum(), 0),
                "tch_estimado_prom": round(est_agg, 1),
                "tch_real_prom": round(real_agg, 1) if wr.sum() else None,
                "n_con_real": len(con_real)})
resumen = pd.DataFrame(res)

# ---- escribir Excel formateado ----
with pd.ExcelWriter(OUT, engine="openpyxl") as xl:
    df.to_excel(xl, sheet_name="Estimados", index=False)
    resumen.to_excel(xl, sheet_name="Resumen por año", index=False)

wb = openpyxl.load_workbook(OUT)
hdr_fill = PatternFill("solid", fgColor="1F5A3D")
hdr_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="D9D9D9")
for ws in wb.worksheets:
    for c in ws[1]:
        c.fill = hdr_fill; c.font = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        wd = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(wd + 2, 10), 22)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = Border(bottom=thin)
wb.save(OUT)
print(f"\nExcel: {OUT}  ({len(df)} filas, años {sorted(df['anio'].unique())})")
