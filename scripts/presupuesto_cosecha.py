# -*- coding: utf-8 -*-
"""Presupuesto de cosecha por suerte y por mes, en Excel con fórmulas.

Suertes: maestro oficial (maestro-riopaila), uso CAÑA de la planta, sin las que están en RENOVACION.
  · próximo corte en el año del presupuesto → esa fecha («Maestro <año>»).
  · próximo corte todavía en el año anterior (faltan por cosechar) → esa fecha + «meses de ciclo»
    («Pendiente <año-1>»). Solo entra la suerte cuya fecha cae en el año del presupuesto.
TCH (en este orden):
  1. Motor satelital (estimados.json de la última corrida), si el estimado es del MISMO corte.
  2. Histórico de la suerte: mediana de sus últimos 3 cortes desde la última siembra
     (ZQM + cosechas_recientes.xlsx, igual que la corrida mensual).
  3. Respaldos: mediana de los cortes de la planta (últimos 6 años) por zona agro × variedad × soca,
     zona × variedad, variedad × soca, soca y planta.
El ajuste (celda editable) se aplica al TCH histórico y de respaldo, no al del motor. La hoja «Supuestos»
trae la prueba de cómo le habría ido al histórico presupuestando el año anterior, para decidir el ajuste.

Uso:
  .venv-gee\\Scripts\\python.exe scripts\\presupuesto_cosecha.py --anio 2027 --planta riopaila \\
      --meses-ciclo 12 --ajuste -0.15 [--salida archivo.xlsx]
Salida por defecto: <TCH_DOCS>/01 Estimados y TCH/Presupuesto_cosecha_<año>_<Planta>.xlsx
"""
import argparse
import io
import json
import os
import sys
import urllib.request
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from tchmotor import dataset_gee as dg  # noqa: E402
from tchmotor import produccion as pr  # noqa: E402
from tchmotor import rutas  # noqa: E402
from tchmotor.agronomia import cargar_var_map, map_variedad  # noqa: E402

EMPRESAS = {"riopaila": ("RIOP",), "castilla": ("CAST", "CAUC")}
NOMBRE = {"riopaila": "Riopaila", "castilla": "Castilla"}
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre",
         "Noviembre", "Diciembre"]
DIAS_MES = 30.4375
FUENTES = ["Motor satelital", "Histórico de la suerte", "Zona agro × variedad × soca", "Zona agro × variedad",
           "Variedad × soca", "Soca", "Mediana de la planta"]


def fecha(s):
    s = str(s or "").strip()
    for f in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s[:10], f).date()
        except ValueError:
            pass
    return None


def num(s):
    try:
        return float(str(s).strip().replace(",", "."))
    except ValueError:
        return None


def texto_maestro():
    try:
        with urllib.request.urlopen(pr.MAESTRO_URL, timeout=20) as r:
            return r.read().decode("utf-8-sig"), pr.MAESTRO_URL
    except Exception as e:
        return rutas.MAESTRO_LOCAL.read_text(encoding="utf-8-sig"), f"{rutas.MAESTRO_LOCAL} (sin red: {type(e).__name__})"


def leer_maestro(planta):
    texto, fuente = texto_maestro()
    m = pd.read_csv(io.StringIO(texto), sep=";", dtype=str).fillna("")
    m.columns = range(m.shape[1])
    m = m[m[20].str.strip().isin(EMPRESAS[planta]) & (m[7].str.strip().str.upper() == "CAÑA")]
    ma = {}
    for _, x in m.iterrows():
        ma[x[0].strip()] = dict(zona=x[1].strip(), finca=x[2].strip(), area=num(x[3]), zagro=x[5].strip(),
                                var=x[6].strip(), siembra=fecha(x[8]), ult=fecha(x[9]), ncorte=int(num(x[10]) or 0),
                                tch_ppto=num(x[12]), tch_ant=num(x[13]), tecnico=x[18].strip(), prox=fecha(x[19]))
    return ma, fuente


def historia(planta):
    """Ciclos cosechados (ZQM + cosechas_recientes), como en la corrida mensual."""
    c = dg.cortes_zqm(dg.leer_zqm())
    maestro, _ = pr.cargar_maestro()
    rec, _ = pr.cosechas_recientes(pr.leer_cosechas_recientes(), maestro, c)
    c = pd.concat([c, rec], ignore_index=True).drop_duplicates("key")
    c = c[c.planta == planta].copy()
    vm = cargar_var_map()
    return pd.DataFrame({"sec_ste": c.sec_ste.values, "corte": c.corte.astype(int).values,
                         "fecha": pd.to_datetime(c.harvest).dt.date.values, "tch": c.tch.values,
                         "var": [map_variedad(v, vm) or v for v in c.variedad], "anio": c.anio.values}
                        ).sort_values("fecha")


def mediana_suerte(sec, ma, por_suerte, hasta=None):
    """Mediana de los últimos 3 cortes de la suerte desde su siembra (y antes de `hasta`, si se da)."""
    g, alias = por_suerte.get(sec), None
    if (g is None or not len(g)) and sec.startswith("3107-"):  # LA LUISA: el ZQM la trae como 3106-*
        alt = "3106-" + sec[5:]
        if alt not in ma and alt in por_suerte:
            g, alias = por_suerte[alt], alt
    if g is None or not len(g):
        return None, 0, None
    if ma[sec]["siembra"]:
        g = g[g.fecha > ma[sec]["siembra"]]
    if hasta is not None:
        g = g[g.fecha < hasta]
    g = g.tail(3)
    return (float(g.tch.median()), int(len(g)), alias) if len(g) else (None, 0, None)


def backtest(hist, ma, por_suerte, anio):
    """Cómo le habría ido a la mediana histórica presupuestando las cosechas de `anio`."""
    filas = []
    for x in hist[hist.anio == anio].itertuples():
        if x.sec_ste not in ma:
            continue
        pred, n, _ = mediana_suerte(x.sec_ste, ma, por_suerte, hasta=x.fecha)
        if pred is not None:
            filas.append((x.fecha.month, x.tch, pred, ma[x.sec_ste]["area"] or 1))
    if not filas:
        return None
    b = pd.DataFrame(filas, columns=["mes", "real", "pred", "ha"])
    w = lambda d, c: float(np.average(d[c], weights=d.ha))  # noqa: E731
    tramos = {nom: d for nom, d in b.groupby(pd.cut(b.mes, [0, 4, 8, 12], labels=["ene-abr", "may-ago", "sep-dic"]),
                                              observed=True)}
    return {"n": len(b), "real": w(b, "real"), "pred": w(b, "pred"), "sesgo": w(b, "pred") / w(b, "real") - 1,
            "tramos": {k: w(d, "pred") / w(d, "real") - 1 for k, d in tramos.items()}}


def grupo_soca(s):
    return "Plantilla" if s == 1 else "Soca 2-3" if s <= 3 else "Soca 4-6" if s <= 6 else "Soca 7-10" if s <= 10 else "Soca 11+"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--anio", type=int, default=date.today().year + 1)
    ap.add_argument("--planta", choices=list(EMPRESAS), default="riopaila")
    ap.add_argument("--meses-ciclo", type=float, default=12.0, help="para las suertes que faltan por cosechar")
    ap.add_argument("--ajuste", type=float, default=0.0, help="al TCH histórico y de respaldo, p. ej. -0.15")
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    anio, planta, nom = a.anio, a.planta, NOMBRE[a.planta]
    salida = Path(a.salida) if a.salida else rutas.ESTIMADOS_DIR / f"Presupuesto_cosecha_{anio}_{nom}.xlsx"
    hoy = date.today()

    ma, fuente_maestro = leer_maestro(planta)
    renov = sorted(s for s, v in ma.items() if v["var"].upper() == "RENOVACION")
    hist = historia(planta)
    por_suerte = {s: g for s, g in hist.groupby("sec_ste")}
    bt = backtest(hist, ma, por_suerte, anio - 1)

    zagro = {s: v["zagro"] for s, v in ma.items()}
    r = hist[hist.anio >= anio - 6].copy()
    r["zagro"] = r.sec_ste.map(zagro)
    r = r[r.zagro.notna() & (r.zagro != "")]

    def tabla(keys, minn):
        g = r.groupby(keys).tch
        med, n = g.median(), g.size()
        return {k: float(med[k]) for k in med.index if n[k] >= minn}

    t_zvs, t_zv, t_vs, t_s = tabla(["zagro", "var", "corte"], 5), tabla(["zagro", "var"], 8), tabla(["var", "corte"], 8), tabla(["corte"], 10)
    t_pl = float(r.tch.median())

    est = {e["sec_ste"]: e for e in json.load(open(rutas.TABLERO_DATA / "estimados.json", encoding="utf-8"))
           if e["planta"] == planta}
    meta = json.load(open(rutas.TABLERO_DATA / "meta.json", encoding="utf-8"))
    fac = json.load(open(rutas.TABLERO_DATA / "factores.json", encoding="utf-8")).get(planta, {})

    filas = []
    for sec, v in ma.items():
        if sec in renov or not v["prox"] or v["prox"].year > anio:
            continue
        pendiente = v["prox"].year < anio
        soca = v["ncorte"] + 1 if pendiente else v["ncorte"]
        inicio = v["prox"] if pendiente else max([d for d in (v["siembra"], v["ult"]) if d] or [None])
        e = est.get(sec)
        motor = float(e["tch_estimado"]) if (e is not None and not pendiente and int(e["soca"]) == v["ncorte"]) else None
        h, nh, alias = mediana_suerte(sec, ma, por_suerte)
        if motor is not None:
            base, fuente = motor, "Motor satelital"
        elif h is not None:
            base, fuente = h, "Histórico de la suerte"
        else:
            z, var = v["zagro"], v["var"]
            base, fuente = next(((t[k], f) for t, k, f in ((t_zvs, (z, var, soca), FUENTES[2]), (t_zv, (z, var), FUENTES[3]),
                                                            (t_vs, (var, soca), FUENTES[4]), (t_s, soca, FUENTES[5]))
                                 if k in t), (t_pl, FUENTES[6]))
        filas.append(dict(sec=sec, zona=f"Zona {v['zona']}" if v["zona"] else "Sin zona", finca=v["finca"],
                          tecnico=v["tecnico"], var=v["var"], zagro=v["zagro"], soca=soca, area=v["area"], inicio=inicio,
                          origen=f"Pendiente {anio - 1}" if pendiente else f"Maestro {anio}", fbase=v["prox"],
                          base=round(base, 1), fuente=fuente, motor=round(motor, 1) if motor else None,
                          mediana=round(h, 1) if h else None, nh=nh, alias=alias, tch_ppto=v["tch_ppto"], tch_ant=v["tch_ant"]))
    df = pd.DataFrame(filas).sort_values(["fbase", "sec"]).reset_index(drop=True)
    df["f_anio"] = [fb if o.startswith("Maestro") else date.fromordinal(fb.toordinal() + round(a.meses_ciclo * DIAS_MES))
                    for fb, o in zip(df.fbase, df.origen)]
    df["tch"] = np.where(df.fuente == FUENTES[0], df.base, df.base * (1 + a.ajuste))
    df["ton"] = df.tch * df.area
    dA = df[[x.year == anio for x in df.f_anio]]
    mh = df[(df.fuente == FUENTES[0]) & df.mediana.notna()]
    pm = float((mh.motor * mh.area).sum() / mh.area.sum()) if len(mh) else None
    ph = float((mh.mediana * mh.area).sum() / mh.area.sum()) if len(mh) else None

    print(f"{nom} {anio}: {len(dA)} suertes · {dA.area.sum():,.1f} ha · {dA.ton.sum():,.0f} t · "
          f"TCH {dA.ton.sum() / dA.area.sum():.1f} (meses de ciclo {a.meses_ciclo:g}, ajuste {a.ajuste:+.0%})")
    print(dA.groupby("fuente").agg(suertes=("sec", "size"), ha=("area", "sum"), ton=("ton", "sum"))
          .assign(tch=lambda x: x.ton / x.ha).round(1).to_string())
    if bt:
        print(f"prueba {anio - 1}: la mediana histórica se habría pasado {bt['sesgo']:+.1%} ({bt['n']} cortes)")
    if pm:
        print(f"motor vs mediana histórica en {len(mh)} suertes: {pm:.1f} vs {ph:.1f}")

    escribir(df, salida, a, anio, planta, nom, renov, bt, pm, ph, len(mh), meta, fac, fuente_maestro, hoy)
    print("guardado:", salida)


# ======================================================================== Excel
F = "Arial"
FB, FN, FH = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=10, bold=True, color="FFFFFF")
FIN, FNOTE, FTIT = Font(name=F, size=10, color="0000FF"), Font(name=F, size=9, italic=True, color="595959"), Font(name=F, size=14, bold=True)
HEAD, TOT, REF = PatternFill("solid", fgColor="2F3A33"), PatternFill("solid", fgColor="F4F4F0"), PatternFill("solid", fgColor="EDEDED")
FILLS_ZONA = [PatternFill("solid", fgColor=c) for c in ("DCE9F8", "E2F0D9", "FFF2CC", "EAD1DC")]
FILL_PLANTA = PatternFill("solid", fgColor="FBE3D6")
YEL = PatternFill("solid", fgColor="FFFF00")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
FUENTE_FILL = {"Motor satelital": "D9EAD3", "Histórico de la suerte": "FFF2CC"}
PCT = "+0.0%;-0.0%;0.0%"


def coma(x, d=1):
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def escribir(df, salida, a, anio, planta, nom, renov, bt, pm, ph, nmh, meta, fac, fuente_maestro, hoy):
    wb = Workbook()
    # ------------------------------------------------------------ Supuestos
    sp = wb.active
    sp.title = "Supuestos"
    sp["A1"] = "Supuestos (celdas amarillas: se pueden cambiar y todo el libro se recalcula)"
    sp["A1"].font = FTIT
    filas = [("Meses de ciclo para las suertes que faltan por cosechar", a.meses_ciclo, "0.0",
              f"Fecha de corte {anio} = fecha de corte {anio - 1} del maestro + estos meses (la regla del maestro es 12,5)."),
             ("Ajuste al TCH que no viene del motor (histórico y respaldos)", a.ajuste, PCT,
              "Se aplica al TCH histórico y de respaldo; el del motor no se toca."),
             ("Año del presupuesto", anio, "0", "Solo entra la suerte cuya fecha de corte cae en este año.")]
    for i, (txt, val, fmt, nota) in enumerate(filas, 3):
        sp.cell(i, 1, txt).font = FN
        c = sp.cell(i, 2, val)
        c.fill, c.font, c.number_format = YEL, FIN, fmt
        sp.cell(i, 3, nota).font = FNOTE
    info = [("Corrida del motor", f"{meta.get('fecha_corrida')} (satélite hasta {meta.get('satelite_hasta')}), factor de "
                                  f"calibración {nom} {coma(fac.get('factor', 1), 4)} con {fac.get('cosechas', 0)} cosechas"),
            ("Maestro de suertes", fuente_maestro),
            ("Cortes cerrados", f"ZQM hasta {meta.get('zqm_hasta')} + cosechas_recientes.xlsx (hasta {meta.get('cosechas_hasta')})"),
            ("Fecha de generación", hoy.strftime("%d/%m/%Y"))]
    for i, (k, v) in enumerate(info, 8):
        sp.cell(i, 1, k).font, sp.cell(i, 2, v).font = FN, FB
    sp["A14"] = "Para decidir el ajuste"
    sp["A14"].font = FN
    avisos = []
    if bt:
        tr = " · ".join(f"{k} {v:+.0%}" for k, v in bt["tramos"].items())
        avisos.append(f"Si la mediana de los últimos 3 cortes se hubiera usado para presupuestar {anio - 1}, se habría desviado "
                      f"{bt['sesgo']:+.0%}: TCH {coma(bt['pred'])} contra {coma(bt['real'])} real en {bt['n']} cortes de {nom} ({tr}).")
    if pm:
        avisos.append(f"En las {nmh} suertes que tienen las dos cifras, el motor satelital (lo que ve hoy en la caña del ciclo "
                      f"{anio}) da {coma(pm)} y la mediana histórica {coma(ph)}: el motor está {abs(pm / ph - 1):.0%} "
                      f"{'por debajo' if pm < ph else 'por encima'}.")
    for i, txt in enumerate(avisos, 15):
        c = sp.cell(i, 1, txt)
        c.font, c.alignment = FB, Alignment(wrap_text=True, vertical="top")
        sp.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        sp.row_dimensions[i].height = 30
    sp.column_dimensions["A"].width, sp.column_dimensions["B"].width, sp.column_dimensions["C"].width = 62, 14, 100
    MESES_REF, AJ_REF, ANIO_REF = "Supuestos!$B$3", "Supuestos!$B$4", "Supuestos!$B$5"

    # ------------------------------------------------------------ Por suerte
    ps = wb.create_sheet("Por suerte")
    cols = [("Suerte", 10), ("Zona", 8), ("Hacienda", 18), ("Técnico", 20), ("Variedad", 12), ("Zona agroecológica", 11),
            (f"Soca {anio}", 7), ("Grupo de soca", 10), ("Área neta (ha)", 9), ("Origen de la fecha", 14),
            ("Fecha en el maestro", 11), (f"Inicio del ciclo {anio}", 11), (f"Fecha de corte {anio}", 11), ("Año", 6),
            ("Mes", 5), ("Edad al corte (meses)", 8), ("TCH base", 8), ("Fuente del TCH", 22), (f"TCH {anio}", 8),
            (f"Toneladas {anio}", 11), ("Ref. motor", 8), ("Ref. mediana histórica", 9), ("Cortes en la mediana", 8),
            ("Ref. TCH ppto (maestro)", 9), ("Ref. TCH anterior", 9)]
    K = {n: get_column_letter(i + 1) for i, (n, _) in enumerate(cols)}
    for i, (n, w) in enumerate(cols, 1):
        c = ps.cell(1, i, n)
        c.font, c.alignment, c.fill = FH, CEN, HEAD
        ps.column_dimensions[get_column_letter(i)].width = w
    ps.row_dimensions[1].height = 44
    n1 = len(df) + 1
    azul = {"Área neta (ha)", "TCH base", "Ref. motor", "Ref. mediana histórica", "Ref. TCH ppto (maestro)",
            "Ref. TCH anterior", "Fecha en el maestro", f"Inicio del ciclo {anio}"}
    fmts = {"Área neta (ha)": "0.00", "Fecha en el maestro": "dd/mm/yyyy", f"Inicio del ciclo {anio}": "dd/mm/yyyy",
            f"Fecha de corte {anio}": "dd/mm/yyyy", "Año": "0", "Mes": "0", "Edad al corte (meses)": "0.0", "TCH base": "0.0",
            f"TCH {anio}": "0.0", f"Toneladas {anio}": "#,##0", "Ref. motor": "0.0", "Ref. mediana histórica": "0.0",
            "Ref. TCH ppto (maestro)": "0.0", "Ref. TCH anterior": "0.0"}
    for k, x in enumerate(df.itertuples(index=False), start=2):
        vals = {"Suerte": x.sec, "Zona": x.zona, "Hacienda": x.finca, "Técnico": x.tecnico, "Variedad": x.var,
                "Zona agroecológica": x.zagro, f"Soca {anio}": x.soca, "Grupo de soca": grupo_soca(x.soca),
                "Área neta (ha)": x.area, "Origen de la fecha": x.origen, "Fecha en el maestro": x.fbase,
                f"Inicio del ciclo {anio}": x.inicio, "TCH base": x.base, "Fuente del TCH": x.fuente, "Ref. motor": x.motor,
                "Ref. mediana histórica": x.mediana, "Cortes en la mediana": x.nh or None,
                "Ref. TCH ppto (maestro)": x.tch_ppto, "Ref. TCH anterior": x.tch_ant}
        for n, val in vals.items():
            ps[f"{K[n]}{k}"] = None if isinstance(val, float) and val != val else val
        J, KF, L, M = (f"{K[n]}{k}" for n in ("Origen de la fecha", "Fecha en el maestro", f"Inicio del ciclo {anio}",
                                              f"Fecha de corte {anio}"))
        ps[M] = f'=IF(LEFT({J},7)="Maestro",{KF},{KF}+ROUND({MESES_REF}*{DIAS_MES},0))'
        ps[f"{K['Año']}{k}"] = f"=YEAR({M})"
        ps[f"{K['Mes']}{k}"] = f"=MONTH({M})"
        ps[f"{K['Edad al corte (meses)']}{k}"] = f'=IF(ISNUMBER({L}),({M}-{L})/{DIAS_MES},"")'
        Q, RF = f"{K['TCH base']}{k}", f"{K['Fuente del TCH']}{k}"
        ps[f"{K[f'TCH {anio}']}{k}"] = f'=IF({RF}="Motor satelital",{Q},{Q}*(1+{AJ_REF}))'
        ps[f"{K[f'Toneladas {anio}']}{k}"] = f"={K[f'TCH {anio}']}{k}*{K['Área neta (ha)']}{k}"
        for n, _ in cols:
            cell = ps[f"{K[n]}{k}"]
            cell.font = FIN if n in azul else FB
            if n in fmts:
                cell.number_format = fmts[n]
        ps[RF].fill = PatternFill("solid", fgColor=FUENTE_FILL.get(x.fuente, "FCE5CD"))
        if isinstance(x.alias, str):
            ps[f"{K['Ref. mediana histórica']}{k}"].comment = Comment(f"Historia tomada del ZQM como {x.alias} (LA LUISA)", "Presupuesto")
    ps.freeze_panes = "B2"
    ps.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n1}"
    ps.cell(n1 + 2, 1, "Azul: dato de la fuente. Negro: fórmula. Verde: TCH del motor; amarillo: histórico de la suerte; "
                       "naranja: respaldo por grupo. «Ref.»: columnas solo de referencia.").font = FNOTE

    def R(n):
        return f"'Por suerte'!${K[n]}$2:${K[n]}${n1}"

    an = f'{R("Año")},{ANIO_REF}'
    area, tons = R("Área neta (ha)"), R(f"Toneladas {anio}")

    # ------------------------------------------------------------ Presupuesto mensual
    pm_ = wb.create_sheet("Presupuesto mensual", 0)
    pm_["A1"] = f"{nom} · Presupuesto de cosecha {anio}"
    pm_["A1"].font = FTIT
    pm_["A2"] = (f"Generado el {hoy.strftime('%d/%m/%Y')}. TCH = toneladas ÷ área. Toneladas = TCH {anio} × área neta de cada "
                 "suerte (hoja «Por suerte»). Supuestos editables en la hoja «Supuestos».")
    pm_["A2"].font = FNOTE
    zonas = sorted(z for z in df.zona.unique())
    grupos = [(z, FILLS_ZONA[i % len(FILLS_ZONA)], z) for i, z in enumerate(zonas)] + [(nom, FILL_PLANTA, None)]
    tot_c0 = 2 + 4 * (len(grupos) - 1)
    for i, (gnom, fill, _) in enumerate(grupos):
        c0 = 2 + 4 * i
        pm_.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 3)
        pm_.cell(4, c0, gnom).font, pm_.cell(4, c0).alignment = FN, CEN
        for j in range(4):
            pm_.cell(4, c0 + j).fill = fill
    heads = ["Mes"] + ["Suertes", "Área (ha)", "Toneladas", "TCH"] * len(grupos)
    ref = None
    if planta == "riopaila" and rutas.PPTO_REFERENCIA.exists():
        ref = json.load(open(rutas.PPTO_REFERENCIA, encoding="utf-8")).get("meses", {})
        ref_c0 = 2 + 4 * len(grupos)
        pm_.merge_cells(start_row=4, start_column=ref_c0, end_row=4, end_column=ref_c0 + 3)
        pm_.cell(4, ref_c0, f"Referencia: presupuesto {anio - 1} de Planeación").font = FN
        pm_.cell(4, ref_c0).alignment = CEN
        for j in range(4):
            pm_.cell(4, ref_c0 + j).fill = REF
        heads += ["Área (ha)", "Toneladas", "TCH", f"Toneladas {anio} vs {anio - 1}"]
    for i, h in enumerate(heads, 1):
        c = pm_.cell(5, i, h)
        c.font, c.alignment, c.fill = FH, CEN, HEAD
    pm_.row_dimensions[5].height = 32
    kt = 18
    L = get_column_letter
    for mi, mnom in enumerate(MESES, 1):
        k = 5 + mi
        pm_.cell(k, 1, mnom).font = FN
        for i, (_, _, zfil) in enumerate(grupos):
            c0 = 2 + 4 * i
            crit = f"{an},{R('Mes')},{mi}" + (f',{R("Zona")},"{zfil}"' if zfil else "")
            pm_[f"{L(c0)}{k}"] = f"=COUNTIFS({crit})"
            pm_[f"{L(c0 + 1)}{k}"] = f"=SUMIFS({area},{crit})"
            pm_[f"{L(c0 + 2)}{k}"] = f"=SUMIFS({tons},{crit})"
            pm_[f"{L(c0 + 3)}{k}"] = f'=IF({L(c0 + 1)}{k}=0,"",{L(c0 + 2)}{k}/{L(c0 + 1)}{k})'
            for j, fmt in enumerate(("0", "#,##0.0", "#,##0", "0.0")):
                pm_[f"{L(c0 + j)}{k}"].number_format, pm_[f"{L(c0 + j)}{k}"].font = fmt, FB
        if ref is not None:
            p = ref.get(f"{anio - 1}-{mi:02d}", {}).get("empresa", {}).get("ppto") or {}
            ca, ct, cc, cv = (L(ref_c0 + j) for j in range(4))
            pm_[f"{ca}{k}"], pm_[f"{ct}{k}"] = p.get("area"), p.get("ton")
            pm_[f"{cc}{k}"] = f'=IF(N({ca}{k})=0,"",{ct}{k}/{ca}{k})'
            pm_[f"{cv}{k}"] = f'=IF(N({ct}{k})=0,"",{L(tot_c0 + 2)}{k}/{ct}{k}-1)'
            for col, fmt, fnt in ((ca, "#,##0.0", FIN), (ct, "#,##0", FIN), (cc, "0.0", FB), (cv, PCT, FB)):
                pm_[f"{col}{k}"].number_format, pm_[f"{col}{k}"].font = fmt, fnt
    pm_.cell(kt, 1, f"Total {anio}").font = FN
    ultima = 1 + len(heads) - 1
    for i in range(len(grupos)):
        c0 = 2 + 4 * i
        for j in range(3):
            pm_[f"{L(c0 + j)}{kt}"] = f"=SUM({L(c0 + j)}6:{L(c0 + j)}17)"
        pm_[f"{L(c0 + 3)}{kt}"] = f'=IF({L(c0 + 1)}{kt}=0,"",{L(c0 + 2)}{kt}/{L(c0 + 1)}{kt})'
        for j, fmt in enumerate(("0", "#,##0.0", "#,##0", "0.0")):
            pm_[f"{L(c0 + j)}{kt}"].number_format = fmt
    if ref is not None:
        ca, ct, cc, cv = (L(ref_c0 + j) for j in range(4))
        pm_[f"{ca}{kt}"], pm_[f"{ct}{kt}"] = f"=SUM({ca}6:{ca}17)", f"=SUM({ct}6:{ct}17)"
        pm_[f"{cc}{kt}"], pm_[f"{cv}{kt}"] = f"={ct}{kt}/{ca}{kt}", f"={L(tot_c0 + 2)}{kt}/{ct}{kt}-1"
        for col, fmt in ((ca, "#,##0.0"), (ct, "#,##0"), (cc, "0.0"), (cv, PCT)):
            pm_[f"{col}{kt}"].number_format = fmt
        pm_.cell(kt + 1, 1, f"Referencia {anio - 1}: presupuesto anual de Planeación (archivo de la app de Turnos). "
                            "El TCH de la referencia es toneladas ÷ área.").font = FNOTE
    for col in range(1, ultima + 1):
        pm_.cell(kt, col).fill, pm_.cell(kt, col).font = TOT, FN
    tot_area, tot_ton = f"$" + L(tot_c0 + 1) + f"${kt}", f"$" + L(tot_c0 + 2) + f"${kt}"
    f0 = kt + 3
    pm_.cell(f0, 1, f"De dónde sale el TCH (suertes del presupuesto {anio})").font = FN
    for i, h in enumerate(["Fuente del TCH", "Suertes", "Área (ha)", "Toneladas", "TCH", "% del área"], 1):
        c = pm_.cell(f0 + 1, i, h)
        c.font, c.alignment, c.fill = FH, CEN, HEAD
    for j, fu in enumerate(FUENTES):
        k = f0 + 2 + j
        crit = f'{an},{R("Fuente del TCH")},"{fu}"'
        pm_.cell(k, 1, fu).font = FB
        pm_[f"B{k}"], pm_[f"C{k}"], pm_[f"D{k}"] = f"=COUNTIFS({crit})", f"=SUMIFS({area},{crit})", f"=SUMIFS({tons},{crit})"
        pm_[f"E{k}"], pm_[f"F{k}"] = f'=IF(C{k}=0,"",D{k}/C{k})', f'=IF({tot_area}=0,"",C{k}/{tot_area})'
        for col, fmt in zip("BCDEF", ("0", "#,##0.0", "#,##0", "0.0", "0.0%")):
            pm_[f"{col}{k}"].number_format, pm_[f"{col}{k}"].font = fmt, FB
    pm_.column_dimensions["A"].width = 28
    for i in range(2, ultima + 1):
        pm_.column_dimensions[L(i)].width = 10
    pm_.freeze_panes = "B6"

    # ------------------------------------------------------------ Variedad y soca
    vs = wb.create_sheet("Variedad y soca", 1)
    vs["A1"] = f"Presupuesto {anio} por variedad y por soca"
    vs["A1"].font = FTIT

    def bloque(fila, titulo, campo, valores):
        vs.cell(fila, 1, titulo).font = FN
        for i, h in enumerate([campo, "Suertes", "Área (ha)", "Toneladas", "TCH", "% de las toneladas"], 1):
            c = vs.cell(fila + 1, i, h)
            c.font, c.alignment, c.fill = FH, CEN, HEAD
        k = fila + 2
        for val in valores:
            crit = f'{an},{R(campo)},"{val}"'
            vs.cell(k, 1, val).font = FB
            vs[f"B{k}"], vs[f"C{k}"], vs[f"D{k}"] = f"=COUNTIFS({crit})", f"=SUMIFS({area},{crit})", f"=SUMIFS({tons},{crit})"
            vs[f"E{k}"] = f'=IF(C{k}=0,"",D{k}/C{k})'
            vs[f"F{k}"] = f"=IF('Presupuesto mensual'!{tot_ton}=0,\"\",D{k}/'Presupuesto mensual'!{tot_ton})"
            for col, fmt in zip("BCDEF", ("0", "#,##0.0", "#,##0", "0.0", "0.0%")):
                vs[f"{col}{k}"].number_format, vs[f"{col}{k}"].font = fmt, FB
            k += 1
        return k

    dA = df[[x.year == anio for x in df.f_anio]]
    variedades = dA.groupby("var").area.sum().sort_values(ascending=False).index.tolist()
    nxt = bloque(3, "Por variedad", "Variedad", variedades)
    bloque(nxt + 1, "Por soca", "Grupo de soca", ["Plantilla", "Soca 2-3", "Soca 4-6", "Soca 7-10", "Soca 11+"])
    vs.column_dimensions["A"].width = 18
    for col in "BCDEF":
        vs.column_dimensions[col].width = 12

    # ------------------------------------------------------------ Cómo se hizo
    ch = wb.create_sheet("Cómo se hizo")
    ch.column_dimensions["A"].width = 125
    txt = [
        ("Qué suertes entran", True),
        (f"Suertes de {nom} con uso CAÑA en el maestro, sin las que están en renovación ({', '.join(renov) or 'ninguna'}). "
         f"Las que ya tienen próximo corte en {anio} entran con esa fecha. Las que todavía tienen próximo corte en {anio - 1} "
         f"(faltan por cosechar) entran con esa fecha más los meses de ciclo de la hoja Supuestos. Solo cuentan las suertes "
         f"cuya fecha de corte cae en {anio}.", False),
        (f"Soca {anio}: para las que ya tienen fecha {anio}, el número de corte del maestro; para las pendientes, ese número + 1.", False),
        ("", False),
        ("De dónde sale el TCH (en este orden)", True),
        (f"1. Motor satelital (corrida del {meta.get('fecha_corrida')}, con su factor de calibración "
         f"{coma(fac.get('factor', 1), 4)}): solo cuando el estimado es del mismo corte que se cosecha y la caña ya tiene "
         "4 meses o más.", False),
        ("2. Histórico de la suerte: mediana de sus últimos 3 cortes desde la última siembra (ZQM y cosechas_recientes.xlsx, "
         "los mismos archivos que usa la corrida mensual). Las suertes 3107 de LA LUISA se buscan como 3106 en el ZQM cuando no "
         "tienen historia propia.", False),
        (f"3. Respaldos, para suertes sin historia desde su siembra: mediana de los cortes de {nom} de los últimos 6 años por zona "
         "agroecológica × variedad × soca (mínimo 5 cortes), zona × variedad (8), variedad × soca (8), soca (10) y, si nada "
         "aplica, la mediana de la planta.", False),
        ("El ajuste de la hoja Supuestos se aplica solo al TCH histórico y de respaldo.", False),
        ("", False),
        ("Límites", True),
        ("El reparto por mes depende de cuándo se corten de verdad las suertes pendientes y de las fechas de Planeación en el "
         "maestro. Conviene volver a generarlo cada mes: a medida que la caña crece, más suertes pasan del histórico al motor.", False),
        ("El área es la neta del maestro; no descuenta área para semilla ni la que no se coseche. No incluye siembras nuevas.", False),
        ("Generado con tch-motor/scripts/presupuesto_cosecha.py.", False),
    ]
    for i, (t, neg) in enumerate(txt, 1):
        c = ch.cell(i, 1, t)
        c.font, c.alignment = (FN if neg else FB), Alignment(wrap_text=True, vertical="top")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title == "Por suerte"
    wb.active = 0
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)


if __name__ == "__main__":
    main()
