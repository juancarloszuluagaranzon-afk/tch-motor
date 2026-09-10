"""Carga los cortes YA COSECHADOS de una campaña a tch_real, desde el ZQM.

Fuente: C:\\Users\\Agr349\\Documents\\historico zqm.xlsx (SAP). Lee SOLO las columnas
no sensibles (nada de liquidación / $). El TCH real correcto = columna `TCH`
(= TC-Total/área); por suerte×año se toma la fila de mayor TC-Total (el corte
completo, no un mes parcial). La variedad V## se mapea a Cenicaña con tch_var_map.

Uso:
  python scripts/cargar_reales.py --campana 2026
  python scripts/cargar_reales.py --campana 2026 --dry
"""
import argparse
import io
import os
import sys
from datetime import date

import openpyxl

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
from tchmotor import Supabase  # noqa

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ZQM = r"C:\Users\Agr349\Documents\historico zqm.xlsx"
# Índices de columna (evita el mojibake de encabezados con acentos).
C_ANIO, C_SEC, C_MES, C_CORTE, C_VAR = 1, 3, 4, 15, 16
C_AREA, C_TCTOT, C_TCH, C_PLANTA = 17, 20, 21, 46


def _env(path, name):
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if line.startswith(name + "="):
            return line.split("=", 1)[1].strip().strip('"').strip("'")
    return None


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _planta(v):
    s = str(v or "").strip().lower()
    if s.startswith("cas"):   # 'CAS' en el ZQM
        return "castilla"
    if s.startswith("rio"):   # 'RIO'
        return "riopaila"
    return s or None


def leer_zqm(campana):
    """{sec_ste: fila} con la fila de mayor TC-Total por suerte de la campaña."""
    wb = openpyxl.load_workbook(ZQM, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    mejor = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[C_ANIO]) != str(campana):
            continue
        sec = row[C_SEC]
        tch = _num(row[C_TCH])
        area = _num(row[C_AREA])
        tctot = _num(row[C_TCTOT])
        if not sec or tch is None or not area:
            continue
        prev = mejor.get(sec)
        if prev is None or (tctot or 0) > (prev["_tctot"] or 0):
            mejor[sec] = {
                "sec_ste": str(sec).strip(),
                "campana": str(campana),
                "planta": _planta(row[C_PLANTA]),
                "corte": int(row[C_CORTE]) if _num(row[C_CORTE]) is not None else None,
                "variedad": str(row[C_VAR]).strip() if row[C_VAR] else None,
                "area_ha": round(area, 4),
                "tc_total": round(tctot, 2) if tctot is not None else None,
                "tch_real": round(tch, 2),
                "fecha_corte_real": None,  # ZQM no trae día; se deja nulo
                "fuente": "zqm",
                "_tctot": tctot,
                "_mes": _num(row[C_MES]),
            }
    wb.close()
    return mejor


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--campana", default=str(date.today().year))
    ap.add_argument("--dry", action="store_true")
    args = ap.parse_args()

    os.environ.setdefault("TCH_SUPABASE_URL", _env(r"C:\Users\Agr349\agrocontrol-campo\.env.local", "NEXT_PUBLIC_SUPABASE_URL"))
    os.environ.setdefault("TCH_SUPABASE_KEY", _env(r"C:\Users\Agr349\turnos-jornales\.env.local", "RIOMAP_SERVICE_ROLE_KEY"))

    filas = list(leer_zqm(args.campana).values())
    # mapear variedad V## → Cenicaña si se puede
    try:
        vm = Supabase().cargar_var_map()
        for f in filas:
            if f["variedad"] in vm:
                f["variedad"] = vm[f["variedad"]]
    except Exception as e:
        print("aviso: sin var_map, dejo la variedad cruda:", str(e)[:80])

    for f in filas:
        f.pop("_tctot", None)
        f.pop("_mes", None)

    print(f"cortes cerrados de {args.campana} en el ZQM: {len(filas)}")
    for f in filas[:8]:
        print(f"  {f['sec_ste']} soca {f['corte']} {f['variedad']} "
              f"[{f['planta']}]: TCH {f['tch_real']} · {f['area_ha']} ha")
    if args.dry:
        print("\n[dry] no se escribió tch_real.")
        return
    if not filas:
        print("nada que cargar.")
        return
    n = Supabase().guardar_reales(filas)
    print(f"\nescritas {n} filas en tch_real (campaña {args.campana}).")


if __name__ == "__main__":
    main()
