# -*- coding: utf-8 -*-
"""Cómo va la caña en pie con 4 meses o más (la que estima el motor), frente al presupuesto y al corte anterior.

Fuente: estimados.json de la última corrida del motor (las suertes de 4 meses o más con satélite suficiente),
cruzado con el maestro oficial (zona, hacienda, técnico, TCH ppto, TCH anterior, próximo corte) y, cuando ya
existe, con el aforo de Agronomía del ciclo en curso. Excel con fórmulas; las comparaciones contra ppto y contra
el corte anterior se ponderan por área y solo con las suertes que tienen ese dato.

Uso:  .venv-gee\\Scripts\\python.exe scripts\\cana_en_pie.py [--planta riopaila] [--salida x.xlsx]
"""
import argparse
import json
import os
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from tchmotor import rutas  # noqa: E402
from presupuesto_cosecha import NOMBRE, leer_maestro  # noqa: E402

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre",
         "Noviembre", "Diciembre"]
EDADES = [("4-6 meses", 0, 6), ("6-9 meses", 6, 9), ("9-12 meses", 9, 12), ("Más de 12", 12, 99)]
SOCAS = [("Plantilla", 1, 1), ("Soca 2-3", 2, 3), ("Soca 4-6", 4, 6), ("Soca 7-10", 7, 10), ("Soca 11+", 11, 99)]

F = "Arial"
FB, FN, FH = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=10, bold=True, color="FFFFFF")
FIN, FNOTE, FTIT = Font(name=F, size=10, color="0000FF"), Font(name=F, size=9, italic=True, color="595959"), Font(name=F, size=14, bold=True)
F9 = Font(name=F, size=9, color="595959")
HEAD, CALC, TOT = (PatternFill("solid", fgColor=c) for c in ("2F3A33", "EFEFEF", "F4F4F0"))
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
PCT = "+0.0%;-0.0%;0.0%"


def banda(v, tabla):
    return next(n for n, lo, hi in tabla if lo < v <= hi) if tabla is EDADES else next(n for n, lo, hi in tabla if lo <= v <= hi)


def armar(planta):
    ma, fuente_maestro = leer_maestro(planta)
    e = pd.DataFrame(json.load(open(rutas.TABLERO_DATA / "estimados.json", encoding="utf-8")))
    e = e[(e.planta == planta) & e.area_ha.notna()].copy()
    try:
        import comparar_aforo_motor_pendientes as cap
        af = cap.leer_aforos()
    except Exception as ex:  # el libro de aforos es opcional
        print("aviso: sin aforos de Agronomía:", type(ex).__name__)
        af = None
    filas = []
    for x in e.itertuples():
        v = ma.get(x.sec_ste, {})
        aforo = f_aforo = None
        if af is not None and v:
            ini = max([d for d in (v.get("siembra"), v.get("ult")) if d] or [None])
            g = af[(af.sec == x.sec_ste) & (af.corte_af == x.soca)]
            if ini:
                g = g[g.f_eval.dt.date > ini]
            if len(g):
                g = g.sort_values("f_eval").iloc[-1]
                aforo, f_aforo = float(g.aforo), g.f_eval.date()
        prox = v.get("prox")
        filas.append(dict(sec=x.sec_ste, finca=v.get("finca"), zona=f"Zona {v['zona']}" if v.get("zona") else "Sin zona",
                          tecnico=v.get("tecnico"), variedad=x.variedad, soca=int(x.soca), edad=float(x.edad_meses),
                          prox=prox, mes=f"{MESES[prox.month - 1]} {prox.year}" if prox else "Sin fecha",
                          orden=(prox.year * 12 + prox.month) if prox else 99999, area=float(x.area_ha),
                          tch=float(x.tch_estimado), std=float(x.tch_std), ppto=v.get("tch_ppto"), ant=v.get("tch_ant"),
                          aforo=aforo, f_aforo=f_aforo))
    d = pd.DataFrame(filas).sort_values(["orden", "sec"]).reset_index(drop=True)
    return d, fuente_maestro


def escribir(d, salida, nom, meta, fac, fuente_maestro):
    wb = Workbook()
    ws = wb.active
    ws.title = "Por suerte"
    cols = [("Suerte", 10, None), ("Hacienda", 18, None), ("Zona", 8, None), ("Técnico", 20, None), ("Variedad", 12, None),
            ("Corte", 6, "0"), ("Grupo de corte", 10, None), ("Edad (meses)", 8, "0.0"), ("Grupo de edad", 10, None),
            ("Corte previsto (maestro)", 11, "dd/mm/yyyy"), ("Mes previsto", 14, None), ("Área neta (ha)", 9, "0.00"),
            ("TCH estimado", 9, "0.0"), ("± TCH", 7, "0.0"), ("Toneladas estimadas", 11, "#,##0"),
            ("TCH ppto (maestro)", 9, "0.0"), ("TCH anterior", 9, "0.0"), ("Frente al ppto", 9, PCT),
            ("Frente al corte anterior", 9, PCT), ("Toneladas de menos vs anterior", 11, "#,##0"),
            ("TCH aforo del ciclo", 9, "0.0"), ("Fecha del aforo", 11, "dd/mm/yyyy"), ("Aforo frente al motor", 9, PCT),
            ("Est × ha (con ppto)", 10, "#,##0"), ("Ppto × ha", 10, "#,##0"), ("Est × ha (con anterior)", 10, "#,##0"),
            ("Anterior × ha", 10, "#,##0"), ("± × ha", 10, "#,##0"), ("Menos de −20 %", 7, "0"), ("−20 a −10 %", 7, "0"),
            ("Dentro de ±10 %", 7, "0"), ("Más de +10 %", 7, "0")]
    K = {n: get_column_letter(i + 1) for i, (n, _, _) in enumerate(cols)}
    for i, (n, w, _) in enumerate(cols, 1):
        c = ws.cell(1, i, n)
        c.alignment = CEN
        c.fill, c.font = (CALC, Font(name=F, size=9, bold=True, color="404040")) if i > 23 else (HEAD, FH)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 44
    n1 = len(d) + 1
    azul = {"Área neta (ha)", "TCH estimado", "± TCH", "TCH ppto (maestro)", "TCH anterior", "TCH aforo del ciclo"}
    nan = lambda v: None if v is None or (isinstance(v, float) and v != v) else v  # noqa: E731
    for k, x in enumerate(d.itertuples(index=False), start=2):
        vals = [x.sec, x.finca, x.zona, x.tecnico, x.variedad, x.soca, banda(x.soca, SOCAS), x.edad, banda(x.edad, EDADES),
                x.prox, x.mes, x.area, x.tch, x.std, None, nan(x.ppto) if (x.ppto or 0) > 0 else None,
                nan(x.ant) if (x.ant or 0) > 0 else None, None, None, None, nan(x.aforo),
                x.f_aforo if isinstance(x.f_aforo, date) else None]
        for j, v in enumerate(vals, 1):
            ws.cell(k, j, nan(v))
        L, M, P, Q, U = (f"{K[n]}{k}" for n in ("Área neta (ha)", "TCH estimado", "TCH ppto (maestro)", "TCH anterior",
                                               "TCH aforo del ciclo"))
        S = f"{K['Frente al corte anterior']}{k}"
        f = {"Toneladas estimadas": f"={M}*{L}",
             "Frente al ppto": f'=IF(ISNUMBER({P}),{M}/{P}-1,"")',
             "Frente al corte anterior": f'=IF(ISNUMBER({Q}),{M}/{Q}-1,"")',
             "Toneladas de menos vs anterior": f'=IF(ISNUMBER({Q}),({Q}-{M})*{L},"")',
             "Aforo frente al motor": f'=IF(ISNUMBER({U}),{U}/{M}-1,"")',
             "Est × ha (con ppto)": f"=IF(ISNUMBER({P}),{M}*{L},0)", "Ppto × ha": f"=IF(ISNUMBER({P}),{P}*{L},0)",
             "Est × ha (con anterior)": f"=IF(ISNUMBER({Q}),{M}*{L},0)", "Anterior × ha": f"=IF(ISNUMBER({Q}),{Q}*{L},0)",
             "± × ha": f"={K['± TCH']}{k}*{L}",
             "Menos de −20 %": f'=IF(ISNUMBER({Q}),IF({S}<-0.2,1,0),0)',
             "−20 a −10 %": f'=IF(ISNUMBER({Q}),IF(AND({S}>=-0.2,{S}<-0.1),1,0),0)',
             "Dentro de ±10 %": f'=IF(ISNUMBER({Q}),IF(ABS({S})<=0.1,1,0),0)',
             "Más de +10 %": f'=IF(ISNUMBER({Q}),IF({S}>0.1,1,0),0)'}
        for n, frm in f.items():
            ws[f"{K[n]}{k}"] = frm
        for j, (n, _, fmt) in enumerate(cols, 1):
            c = ws.cell(k, j)
            c.font = F9 if j > 23 else (FIN if n in azul and c.value is not None else FB)
            if fmt:
                c.number_format = fmt
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n1}"
    ws.cell(n1 + 2, 1, "Azul: dato de la fuente (corrida del motor, maestro, libro de aforos). Negro: fórmula. Gris: columnas "
                       "de cálculo del resumen.").font = FNOTE

    def R(n):
        return f"'Por suerte'!${K[n]}$2:${K[n]}${n1}"

    rs = wb.create_sheet("Resumen", 0)
    rs["A1"] = f"{nom} · caña en pie con 4 meses o más"
    rs["A1"].font = FTIT
    rs["A2"] = (f"Corrida del motor del {meta.get('fecha_corrida')} (satélite hasta {meta.get('satelite_hasta')}, factor de "
                f"calibración {str(fac.get('factor', 1)).replace('.', ',')}). TCH ponderado por área. «Frente al ppto» y «frente al "
                "corte anterior» comparan solo las suertes que tienen ese dato en el maestro.")
    rs["A2"].font, rs["A2"].alignment = FNOTE, Alignment(wrap_text=True)
    rs.merge_cells("A2:L2")
    rs.row_dimensions[2].height = 30
    heads = ["Grupo", "Suertes", "Área (ha)", "TCH estimado", "Toneladas estimadas", "± TCH (ponderado)", "Frente al ppto",
             "Frente al corte anterior", "Más de 20 % abajo", "10 a 20 % abajo", "Dentro de ±10 %", "Más de 10 % arriba"]
    fila = 4

    def tabla(titulo, campo, valores):
        nonlocal fila
        rs.cell(fila, 1, titulo).font = FN
        for i, h in enumerate(heads, 1):
            c = rs.cell(fila + 1, i, h)
            c.font, c.alignment, c.fill = FH, CEN, HEAD
        rs.row_dimensions[fila + 1].height = 32
        k = fila + 2
        for val in valores:
            cr = "" if val is None else f',{R(campo)},"{val}"'
            area = f"SUMIFS({R('Área neta (ha)')}{cr})" if cr else f"SUM({R('Área neta (ha)')})"
            s = (lambda col: f"SUMIFS({R(col)}{cr})") if cr else (lambda col: f"SUM({R(col)})")
            rs.cell(k, 1, "Total" if val is None else val).font = FN if val is None else FB
            vals = [f"=COUNTIFS({R('Suerte')},\"<>\"{cr})", f"={area}", f'=IF(C{k}=0,"",{s("Toneladas estimadas")}/C{k})',
                    f"={s('Toneladas estimadas')}", f'=IF(C{k}=0,"",{s("± × ha")}/C{k})',
                    f'=IF({s("Ppto × ha")}=0,"",{s("Est × ha (con ppto)")}/{s("Ppto × ha")}-1)',
                    f'=IF({s("Anterior × ha")}=0,"",{s("Est × ha (con anterior)")}/{s("Anterior × ha")}-1)',
                    f"={s('Menos de −20 %')}", f"={s('−20 a −10 %')}", f"={s('Dentro de ±10 %')}", f"={s('Más de +10 %')}"]
            for i, (x, fm) in enumerate(zip(vals, ["0", "#,##0.0", "0.0", "#,##0", "0.0", PCT, PCT, "0", "0", "0", "0"]), 2):
                c = rs.cell(k, i, x)
                c.number_format, c.font = fm, FN if val is None else FB
            if val is None:
                for col in range(1, 13):
                    rs.cell(k, col).fill = TOT
            k += 1
        fila = k + 1

    tabla("Total", None, [None])
    tabla("Por mes de cosecha prevista", "Mes previsto", list(dict.fromkeys(d.mes)))
    tabla("Por edad", "Grupo de edad", [n for n, _, _ in EDADES if (d.edad.apply(lambda v: banda(v, EDADES)) == n).any()])
    tabla("Por zona", "Zona", sorted(d.zona.unique()))
    tabla("Por corte", "Grupo de corte", [n for n, _, _ in SOCAS if (d.soca.apply(lambda v: banda(v, SOCAS)) == n).any()])
    tabla("Por variedad (de más a menos área)", "Variedad", d.groupby("variedad").area.sum().sort_values(ascending=False).index.tolist())
    tabla("Por hacienda (de más a menos área)", "Hacienda", d.groupby("finca").area.sum().sort_values(ascending=False).index.tolist())
    rs.column_dimensions["A"].width = 24
    for i in range(2, 13):
        rs.column_dimensions[get_column_letter(i)].width = 11

    # ---- Más atrás (foto ordenada)
    ma_ = wb.create_sheet("Más atrás", 1)
    ma_["A1"] = "Suertes que más toneladas pierden frente a su corte anterior"
    ma_["A1"].font = FTIT
    ma_["A2"] = "Foto al generar el libro (no se recalcula). Toneladas de menos = (TCH anterior − TCH estimado) × área."
    ma_["A2"].font = FNOTE
    t = d[d.ant.notna() & (d.ant > 0)].assign(vs=lambda x: x.tch / x.ant - 1, menos=lambda x: (x.ant - x.tch) * x.area)
    t = t[t.menos > 0].sort_values("menos", ascending=False)
    hh = ["Suerte", "Hacienda", "Técnico", "Variedad", "Corte", "Edad (meses)", "Corte previsto", "Área (ha)", "TCH estimado",
          "TCH anterior", "Frente al anterior", "Toneladas de menos", "TCH aforo del ciclo"]
    for i, h in enumerate(hh, 1):
        c = ma_.cell(4, i, h)
        c.font, c.alignment, c.fill = FH, CEN, HEAD
    ma_.row_dimensions[4].height = 32
    for k, x in enumerate(t.itertuples(index=False), start=5):
        vals = [x.sec, x.finca, x.tecnico, x.variedad, x.soca, x.edad, x.prox, x.area, x.tch, x.ant, x.vs, x.menos,
                None if x.aforo is None or x.aforo != x.aforo else x.aforo]
        for i, (v, fm) in enumerate(zip(vals, [None, None, None, None, "0", "0.0", "dd/mm/yyyy", "0.00", "0.0", "0.0", PCT,
                                              "#,##0", "0.0"]), 1):
            c = ma_.cell(k, i, v)
            c.font = FB
            if fm:
                c.number_format = fm
    for i, w in enumerate([10, 18, 20, 12, 6, 8, 11, 8, 9, 9, 9, 10, 9], 1):
        ma_.column_dimensions[get_column_letter(i)].width = w
    ma_.freeze_panes = "B5"

    # ---- Notas
    n = wb.create_sheet("Notas")
    n.column_dimensions["A"].width = 120
    notas = [("Cómo se armó", True),
             (f"Suertes: las que estimó el motor en la corrida del {meta.get('fecha_corrida')}: caña de 4 meses o más con satélite "
              "suficiente. La edad es la de la fecha de la corrida.", False),
             (f"Maestro: {fuente_maestro}. Zona, hacienda, técnico, TCH ppto, TCH anterior y fecha del próximo corte.", False),
             ("TCH ppto: el que trae hoy el maestro. En las suertes que se cortan el año siguiente puede ser todavía el del "
              "presupuesto del año en curso.", False),
             ("TCH anterior: el del corte anterior de la suerte (el ciclo que se cosechó antes del que está en pie).", False),
             ("Aforo del ciclo: «Promedio» del libro «Aforos con Diámetro» de Agronomía, último aforo del mismo corte después "
              "del último corte o la siembra. Hacia los 10 meses; la caña más joven todavía no tiene.", False),
             ("± TCH: dispersión entre los árboles del modelo, ya con el factor de calibración. En la caña de 4 a 9 meses el "
              "estimado se apoya más en la historia de la suerte que en lo que ve el satélite.", False),
             ("Generado con tch-motor/scripts/cana_en_pie.py.", False)]
    for i, (txt, neg) in enumerate(notas, 1):
        c = n.cell(i, 1, txt)
        c.font, c.alignment = (FN if neg else FB), Alignment(wrap_text=True, vertical="top")
    for w_ in wb.worksheets:
        w_.sheet_view.showGridLines = w_.title == "Por suerte"
    wb.active = 0
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--planta", choices=list(NOMBRE), default="riopaila")
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    meta = json.load(open(rutas.TABLERO_DATA / "meta.json", encoding="utf-8"))
    fac = json.load(open(rutas.TABLERO_DATA / "factores.json", encoding="utf-8")).get(a.planta, {})
    nom = NOMBRE[a.planta]
    salida = Path(a.salida) if a.salida else rutas.ESTIMADOS_DIR / f"Cana_en_pie_{nom}_{meta.get('fecha_corrida')}.xlsx"
    d, fuente_maestro = armar(a.planta)
    w = lambda g, c: (g[c] * g.area).sum() / g.area.sum()  # noqa: E731
    gp, ga = d[d.ppto.fillna(0) > 0], d[d.ant.fillna(0) > 0]
    print(f"{nom}: {len(d)} suertes · {d.area.sum():,.1f} ha · TCH {w(d, 'tch'):.1f} · {(d.tch * d.area).sum():,.0f} t · "
          f"vs ppto {w(gp, 'tch') / w(gp, 'ppto') - 1:+.1%} · vs anterior {w(ga, 'tch') / w(ga, 'ant') - 1:+.1%} · "
          f"con aforo del ciclo {int(d.aforo.notna().sum())}")
    escribir(d, salida, nom, meta, fac, fuente_maestro)
    print("guardado:", salida)


if __name__ == "__main__":
    main()
